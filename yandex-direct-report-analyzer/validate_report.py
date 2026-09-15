#!/usr/bin/env python3
"""Проверка Excel-отчёта с расчётами, созданного скриптом build_report.py."""

from __future__ import annotations

import argparse
import json
import math
import posixpath
import re
from pathlib import Path
from xml.etree import ElementTree
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


FONT_NAME = "IBM Plex Sans"
FONT_SIZE = 10
UNDEFINED = {"", "неизвестно", "unknown", "undefined", "not set", "n/a", "n a", "-", "—"}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--analysis", required=True, type=Path)
    parser.add_argument("--workbook", required=True, type=Path)
    return parser.parse_args()


def normalize(value: object) -> str:
    text = "" if value is None else str(value)
    text = text.replace("ё", "е").replace("Ё", "Е").strip().lower()
    text = re.sub(r"[\s\u00a0]+", " ", text)
    return text


def is_undefined(value: object) -> bool:
    text = normalize(value)
    return text in UNDEFINED or text.startswith("не определ") or text.startswith("unknown")


def safe_sheet_name(label: str, existing: set[str]) -> str:
    base = re.sub(r"[\\/*?:\[\]]", " ", label).strip()[:31] or "Срез"
    candidate = base
    suffix = 2
    while candidate in existing:
        tail = f" {suffix}"
        candidate = f"{base[:31 - len(tail)]}{tail}"
        suffix += 1
    existing.add(candidate)
    return candidate


def close(left: float, right: float, tolerance: float = 0.01) -> bool:
    return math.isclose(float(left), float(right), rel_tol=1e-9, abs_tol=tolerance)


def worksheet_members(workbook_path: Path) -> dict[str, str]:
    """Сопоставить видимые имена листов с их XML-файлами внутри книги Excel."""
    spreadsheet_ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    document_rel_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    package_rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    with ZipFile(workbook_path) as archive:
        workbook_root = ElementTree.fromstring(archive.read("xl/workbook.xml"))
        relationship_root = ElementTree.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    relationships = {
        relation.attrib["Id"]: relation.attrib["Target"]
        for relation in relationship_root.findall(f"{{{package_rel_ns}}}Relationship")
    }
    members: dict[str, str] = {}
    sheets = workbook_root.find(f"{{{spreadsheet_ns}}}sheets")
    if sheets is None:
        return members
    for sheet in sheets:
        relation_id = sheet.attrib[f"{{{document_rel_ns}}}id"]
        target = relationships[relation_id]
        if target.startswith("/"):
            member = target.lstrip("/")
        else:
            member = posixpath.normpath(posixpath.join("xl", target))
        members[sheet.attrib["name"]] = member
    return members


def worksheet_xml_flags(workbook_path: Path, members: dict[str, str]) -> dict[str, dict[str, bool]]:
    """Один раз прочитать XML листа, не создавая миллионы оформленных объектов Cell."""
    flags: dict[str, dict[str, bool]] = {}
    with ZipFile(workbook_path) as archive:
        for sheet_name, member in members.items():
            state = {"table": False, "frozen": False, "error": False}
            tail = b""
            with archive.open(member) as source:
                while True:
                    chunk = source.read(1024 * 1024)
                    if not chunk:
                        break
                    block = tail + chunk
                    state["table"] = state["table"] or b"<tablePart" in block
                    state["frozen"] = state["frozen"] or (
                        b'state="frozen"' in block or b'state="frozenSplit"' in block
                    )
                    state["error"] = state["error"] or b't="e"' in block
                    tail = block[-128:]
            flags[sheet_name] = state
    return flags


def main() -> None:
    args = parse_args()
    data = json.loads(args.analysis.read_text(encoding="utf-8"))
    workbook = load_workbook(args.workbook, data_only=False, read_only=True)
    xml_flags = worksheet_xml_flags(args.workbook, worksheet_members(args.workbook))

    existing = {"Резюме", "Методика", "Название кампании"}
    expected_sheets = ["Резюме", "Методика", "Название кампании"]
    slice_names: dict[str, str] = {}
    for slice_data in data.get("slices", []):
        name = safe_sheet_name(slice_data.get("sheet") or slice_data["label"], existing)
        slice_names[slice_data["key"]] = name
        expected_sheets.append(name)
    if workbook.sheetnames != expected_sheets:
        raise SystemExit(f"Unexpected sheets: {workbook.sheetnames}; expected {expected_sheets}")

    required_headers = {
        "Уровень анализа",
        "Название кампании",
        "Сегмент",
        "Расход, ₽",
        "Доля расхода в скопе, %",
        "Показы",
        "Клики",
        "CTR",
        "CPC, ₽",
        "Конверсии",
        "CR",
        "CPL, ₽",
        "Погрешность",
        "KPI CPL всех РК, ₽",
        "Расчёт к всем РК",
        "Статус данных",
    }
    metrics = data.get("metrics_available", {})
    revenue_available = bool(metrics.get("revenue"))
    drr_available = bool(metrics.get("drr"))
    romi_available = bool(metrics.get("romi"))
    aov_available = bool(metrics.get("aov"))
    if revenue_available:
        required_headers.add("Выручка, ₽")
    if drr_available:
        required_headers.update(
            {
                "ДРР",
                "KPI ДРР всех РК",
                "Отклонение ДРР к всем РК, п.п.",
                "Расчёт к всем РК по ДРР",
            }
        )
    if romi_available:
        required_headers.update(
            {
                "ROMI",
                "KPI ROMI всех РК",
                "Отклонение ROMI к всем РК, п.п.",
                "Расчёт к всем РК по ROMI",
            }
        )
    if aov_available:
        required_headers.update(
            {
                "Количество заказов",
                "AOV, ₽",
                "KPI AOV всех РК, ₽",
                "Отклонение AOV к всем РК",
            }
        )
    undefined_segments: list[str] = []
    formula_errors = [name for name, flags in xml_flags.items() if flags["error"]]
    style_errors: list[str] = []
    row_count_errors: list[str] = []
    forbidden_headers: list[str] = []

    expected_rows = {"Название кампании": len(data.get("campaigns", []))}
    for slice_data in data.get("slices", []):
        rows = slice_data.get("rows", [])
        if slice_data.get("key") == "group":
            rows = [row for row in rows if row.get("level") != "Все РК"]
        expected_rows[slice_names[slice_data["key"]]] = len(rows)

    for sheet_name in expected_sheets[2:]:
        sheet = workbook[sheet_name]
        header_cells = next(sheet.iter_rows(min_row=1, max_row=1))
        headers = [cell.value for cell in header_cells]
        missing = required_headers.difference(headers)
        if missing:
            raise SystemExit(f"Missing headers on {sheet_name}: {sorted(missing)}")
        disabled_header_groups = {
            "revenue": {"Выручка, ₽"},
            "drr": {"ДРР", "KPI ДРР всех РК", "Отклонение ДРР к всем РК, п.п."},
            "romi": {"ROMI", "KPI ROMI всех РК", "Отклонение ROMI к всем РК, п.п."},
            "aov": {"AOV, ₽", "KPI AOV всех РК", "Отклонение AOV к всем РК"},
        }
        unexpected = sorted(
            header
            for metric, metric_headers in disabled_header_groups.items()
            if not metrics.get(metric)
            for header in metric_headers
            if header in headers
        )
        if unexpected:
            raise SystemExit(
                f"Disabled metric headers on {sheet_name}: {unexpected}"
            )
        for header in headers:
            normalized = normalize(header)
            if "рекомендац" in normalized or "действ" in normalized:
                forbidden_headers.append(f"{sheet_name}: {header}")
        segment_col = headers.index("Сегмент") + 1
        level_col = headers.index("Уровень анализа") + 1
        query_category_sheet = sheet_name == slice_names.get("query_category")
        query_source_col = None
        campaign_col = None
        query_keys: set[tuple[str, str, str, str]] = set()
        impressions_col = headers.index("Показы") + 1
        clicks_col = headers.index("Клики") + 1
        ctr_col = headers.index("CTR") + 1
        cost_col = headers.index("Расход, ₽") + 1
        cpc_col = headers.index("CPC, ₽") + 1
        if query_category_sheet:
            if "Источник запроса" not in headers:
                raise SystemExit(
                    f"Missing Источник запроса on query category sheet: {sheet_name}"
                )
            query_source_col = headers.index("Источник запроса") + 1
            campaign_col = headers.index("Название кампании") + 1
        if sheet.max_row - 1 != expected_rows[sheet_name]:
            row_count_errors.append(
                f"{sheet_name}: {sheet.max_row - 1} rows; expected {expected_rows[sheet_name]}"
            )
        if not xml_flags.get(sheet_name, {}).get("table"):
            style_errors.append(f"{sheet_name}: no table")
        if not xml_flags.get(sheet_name, {}).get("frozen"):
            style_errors.append(f"{sheet_name}: freeze panes absent")
        relevant_columns = [segment_col, level_col]
        if query_source_col is not None and campaign_col is not None:
            relevant_columns.extend([query_source_col, campaign_col])
        first_col = min(relevant_columns)
        last_col = max(relevant_columns)
        for row_number, values in enumerate(
            sheet.iter_rows(
                min_row=2,
                min_col=first_col,
                max_col=last_col,
                values_only=True,
            ),
            start=2,
        ):
            segment = values[segment_col - first_col]
            if is_undefined(segment):
                undefined_segments.append(f"{sheet_name}!row {row_number}={segment}")
            if sheet_name == slice_names.get("group"):
                level = values[level_col - first_col]
                if level != "Внутри РК":
                    raise SystemExit(
                        f"Group sheet contains non-campaign-local row: {sheet_name}!{row_number}"
                    )
            if query_source_col is not None and campaign_col is not None:
                level = str(values[level_col - first_col] or "")
                campaign = str(values[campaign_col - first_col] or "")
                source = values[query_source_col - first_col]
                if is_undefined(source):
                    raise SystemExit(
                        f"Undefined query source on {sheet_name}!row {row_number}"
                    )
                query_key = (
                    level,
                    normalize(campaign) if level == "Внутри РК" else "",
                    normalize(source),
                    normalize(segment),
                )
                if query_key in query_keys:
                    raise SystemExit(
                        f"Duplicate query-source/category aggregate on "
                        f"{sheet_name}!row {row_number}: {query_key}"
                    )
                query_keys.add(query_key)

        # CTR и CPC должны рассчитываться из агрегированных числителей и знаменателей
        # каждой строки, а при нулевом знаменателе оставаться пустыми.
        first_metric_col = min(impressions_col, clicks_col, ctr_col, cost_col, cpc_col)
        last_metric_col = max(impressions_col, clicks_col, ctr_col, cost_col, cpc_col)
        for row_number, values in enumerate(
            sheet.iter_rows(
                min_row=2,
                min_col=first_metric_col,
                max_col=last_metric_col,
                values_only=True,
            ),
            start=2,
        ):
            impressions = float(values[impressions_col - first_metric_col] or 0)
            clicks = float(values[clicks_col - first_metric_col] or 0)
            cost = float(values[cost_col - first_metric_col] or 0)
            actual_ctr = values[ctr_col - first_metric_col]
            actual_cpc = values[cpc_col - first_metric_col]
            expected_ctr = clicks / impressions if impressions else None
            expected_cpc = cost / clicks if clicks else None
            for label, actual, expected in (
                ("CTR", actual_ctr, expected_ctr),
                ("CPC", actual_cpc, expected_cpc),
            ):
                if actual is None and expected is None:
                    continue
                if actual is None or expected is None or not close(actual, expected, tolerance=1e-9):
                    raise SystemExit(
                        f"{label} mismatch on {sheet_name}!row {row_number}: "
                        f"workbook={actual}, expected={expected}"
                    )

        money = {
            "Расход, ₽",
            "Выручка, ₽",
            "CPC, ₽",
            "CPL, ₽",
            "AOV, ₽",
            "KPI CPL кампании, ₽",
            "KPI CPL всех РК, ₽",
            "KPI AOV кампании, ₽",
            "KPI AOV всех РК, ₽",
        }
        percentages = {
            "Доля расхода в скопе, %",
            "CTR",
            "CR",
            "Погрешность",
            "ДРР",
            "ROMI",
            "KPI ДРР кампании",
            "KPI ДРР всех РК",
            "KPI ROMI кампании",
            "KPI ROMI всех РК",
        }

        def expected_format(header: object) -> str | None:
            if header in money:
                return '#,##0.00 "₽"'
            if header in percentages:
                return "0.00%"
            if isinstance(header, str) and (
                header.startswith("Расчёт") or header.startswith("Отклонение")
            ):
                return "+0.0%;-0.0%;0.0%"
            if isinstance(header, str) and "KPI CPL" in header and header.endswith("×"):
                return '0.00"×"'
            if header in {"Показы", "Клики", "Конверсии", "Количество заказов"}:
                return "#,##0"
            return None

        # Проверяем две соседние строки с данными: предыдущая ошибка применяла оформление только
        # ко второй строке. Также проверяем последнюю строку. Пустые ячейки тоже учитываем,
        # потому что сетка таблицы должна оставаться непрерывной и в необязательных столбцах KPI.
        style_row_groups = [(1, min(sheet.max_row, 3))]
        if sheet.max_row > 3:
            style_row_groups.append((sheet.max_row, sheet.max_row))
        sheet_style_failed = False
        for first_style_row, last_style_row in style_row_groups:
            for row_number, cells in enumerate(
                sheet.iter_rows(min_row=first_style_row, max_row=last_style_row),
                start=first_style_row,
            ):
                for column_number, cell in enumerate(cells, start=1):
                    coordinate = f"{get_column_letter(column_number)}{row_number}"
                    font = getattr(cell, "font", None)
                    if font is None or font.name != FONT_NAME or font.sz != FONT_SIZE:
                        style_errors.append(
                            f"{sheet_name}!{coordinate}: font="
                            f"{getattr(font, 'name', None)} {getattr(font, 'sz', None)}"
                        )
                        sheet_style_failed = True
                        break
                    border = getattr(cell, "border", None)
                    sides = (
                        getattr(border, "left", None),
                        getattr(border, "right", None),
                        getattr(border, "top", None),
                        getattr(border, "bottom", None),
                    )
                    if any(side is None or side.style is None for side in sides):
                        style_errors.append(f"{sheet_name}!{coordinate}: incomplete border")
                        sheet_style_failed = True
                        break
                    required_format = expected_format(headers[column_number - 1])
                    actual_format = getattr(cell, "number_format", "General")
                    if row_number > 1 and required_format and actual_format != required_format:
                        style_errors.append(
                            f"{sheet_name}!{coordinate}: number format={actual_format}; "
                            f"expected {required_format}"
                        )
                        sheet_style_failed = True
                        break
                if sheet_style_failed:
                    break
            if sheet_style_failed:
                break

    summary_sheet = workbook["Резюме"]
    summary_count_labels = {
        "Показы",
        "Клики",
        "Конверсии",
        "Количество заказов",
        "Конверсии как заказы для AOV",
        "Кампаний",
        "Аналитических строк",
        "Устойчивых строк",
        "Гипотез",
        "Аномалий",
        "x3-аномалий",
    }
    for row_number in range(1, summary_sheet.max_row + 1):
        label = summary_sheet.cell(row_number, 1).value
        if label in summary_count_labels:
            actual_format = summary_sheet.cell(row_number, 2).number_format
            if actual_format != "#,##0":
                style_errors.append(
                    f"Резюме!B{row_number}: number format={actual_format}; expected #,##0"
                )
        if label == "Лист":
            for detail_row in range(row_number + 1, summary_sheet.max_row + 1):
                for column in (3, 4):
                    actual_format = summary_sheet.cell(detail_row, column).number_format
                    if actual_format != "#,##0":
                        coordinate = f"{get_column_letter(column)}{detail_row}"
                        style_errors.append(
                            f"Резюме!{coordinate}: number format={actual_format}; expected #,##0"
                        )
            break

    total_impressions = float(data["source_total"].get("impressions") or 0)
    total_clicks = float(data["source_total"].get("clicks") or 0)
    total_cost = float(data["source_total"].get("cost") or 0)
    expected_summary_metrics = {
        "Показы": total_impressions,
        "CTR": total_clicks / total_impressions if total_impressions else None,
        "CPC, ₽": total_cost / total_clicks if total_clicks else None,
    }
    summary_values = {
        summary_sheet.cell(row_number, 1).value: summary_sheet.cell(row_number, 2).value
        for row_number in range(1, summary_sheet.max_row + 1)
    }
    for label, expected in expected_summary_metrics.items():
        actual = summary_values.get(label)
        if actual is None and expected is None:
            continue
        if actual is None or expected is None or not close(actual, expected, tolerance=1e-9):
            raise SystemExit(
                f"Summary {label} mismatch: workbook={actual}, expected={expected}"
            )

    method_sheet = workbook["Методика"]
    method_count_labels = {
        "Малая выборка при CR=100%, кликов",
        "Общие показы",
        "Общие клики",
        "Общие конверсии",
        "Детальных строк",
        "Неактивных строк исключено",
        "Строк с конверсиями > кликов",
    }
    for row_number in range(1, method_sheet.max_row + 1):
        if method_sheet.cell(row_number, 1).value in method_count_labels:
            actual_format = method_sheet.cell(row_number, 2).number_format
            if actual_format != "#,##0":
                style_errors.append(
                    f"Методика!B{row_number}: number format={actual_format}; expected #,##0"
                )

    method_values = {
        method_sheet.cell(row_number, 1).value: method_sheet.cell(row_number, 2).value
        for row_number in range(1, method_sheet.max_row + 1)
    }
    expected_method_metrics = {
        "Общие показы": total_impressions,
        "CTR всех РК": total_clicks / total_impressions if total_impressions else None,
        "CPC всех РК, ₽": total_cost / total_clicks if total_clicks else None,
    }
    for label, expected in expected_method_metrics.items():
        actual = method_values.get(label)
        if actual is None and expected is None:
            continue
        if actual is None or expected is None or not close(actual, expected, tolerance=1e-9):
            raise SystemExit(
                f"Method {label} mismatch: workbook={actual}, expected={expected}"
            )

    campaign_sheet = workbook["Название кампании"]
    campaign_headers = [
        cell.value for cell in next(campaign_sheet.iter_rows(min_row=1, max_row=1))
    ]
    control_columns = {
        "cost": "Расход, ₽",
        "impressions": "Показы",
        "clicks": "Клики",
        "conversions": "Конверсии",
    }
    if revenue_available:
        control_columns["revenue"] = "Выручка, ₽"
    if aov_available:
        control_columns["orders"] = "Количество заказов"
    control_totals = {}
    control_positions = {
        key: campaign_headers.index(header)
        for key, header in control_columns.items()
    }
    control_totals = {key: 0.0 for key in control_columns}
    for values in campaign_sheet.iter_rows(min_row=2, values_only=True):
        for key, position in control_positions.items():
            control_totals[key] += float(values[position] or 0)
    for key, total in control_totals.items():
        expected_total = (
            data["source_total"].get("conversions", 0)
            if key == "orders" and data.get("aov_from_conversions")
            else data["source_total"].get(key, 0)
        )
        if not close(total, expected_total):
            raise SystemExit(
                f"Control total mismatch for {key}: workbook={total}, source={expected_total}"
            )

    if drr_available or romi_available:
        cost_position = campaign_headers.index("Расход, ₽")
        revenue_position = campaign_headers.index("Выручка, ₽")
        total_cost = float(data["source_total"].get("cost") or 0)
        total_revenue = float(data["source_total"].get("revenue") or 0)
        all_drr = total_cost / total_revenue if total_revenue else None
        all_romi = (total_revenue - total_cost) / total_cost if total_cost else None

        def verify_optional_ratio(
            actual: object,
            expected: float | None,
            label: str,
            row_number: int,
        ) -> None:
            if actual is None and expected is None:
                return
            if (
                actual is None
                or expected is None
                or not math.isclose(
                    float(actual),
                    float(expected),
                    rel_tol=1e-9,
                    abs_tol=1e-9,
                )
            ):
                raise SystemExit(
                    f"{label} mismatch on campaign row {row_number}: "
                    f"workbook={actual}, expected={expected}"
                )

        positions: dict[str, int] = {}
        if drr_available:
            positions.update(
                {
                    "drr": campaign_headers.index("ДРР"),
                    "drr_deviation": campaign_headers.index(
                        "Отклонение ДРР к всем РК, п.п."
                    ),
                }
            )
        if romi_available:
            positions.update(
                {
                    "romi": campaign_headers.index("ROMI"),
                    "romi_deviation": campaign_headers.index(
                        "Отклонение ROMI к всем РК, п.п."
                    ),
                }
            )
        for row_number, values in enumerate(
            campaign_sheet.iter_rows(min_row=2, values_only=True),
            start=2,
        ):
            cost = float(values[cost_position] or 0)
            revenue = float(values[revenue_position] or 0)
            if drr_available:
                expected_drr = cost / revenue if revenue else None
                verify_optional_ratio(
                    values[positions["drr"]], expected_drr, "DRR", row_number
                )
                expected_deviation = (
                    expected_drr - all_drr
                    if expected_drr is not None and all_drr is not None
                    else None
                )
                verify_optional_ratio(
                    values[positions["drr_deviation"]],
                    expected_deviation,
                    "DRR deviation",
                    row_number,
                )
            if romi_available:
                expected_romi = (revenue - cost) / cost if cost else None
                verify_optional_ratio(
                    values[positions["romi"]], expected_romi, "ROMI", row_number
                )
                expected_deviation = (
                    expected_romi - all_romi
                    if expected_romi is not None and all_romi is not None
                    else None
                )
                verify_optional_ratio(
                    values[positions["romi_deviation"]],
                    expected_deviation,
                    "ROMI deviation",
                    row_number,
                )

    if aov_available:
        revenue_position = campaign_headers.index("Выручка, ₽")
        orders_position = campaign_headers.index("Количество заказов")
        aov_position = campaign_headers.index("AOV, ₽")
        for row_number, values in enumerate(
            campaign_sheet.iter_rows(min_row=2, values_only=True),
            start=2,
        ):
            revenue = float(values[revenue_position] or 0)
            orders = float(values[orders_position] or 0)
            actual_aov = values[aov_position]
            expected_aov = revenue / orders if orders else None
            if expected_aov is None and actual_aov is None:
                continue
            if expected_aov is None or actual_aov is None or not close(actual_aov, expected_aov):
                raise SystemExit(
                    f"AOV mismatch on campaign row {row_number}: "
                    f"workbook={actual_aov}, expected={expected_aov}"
                )

    if undefined_segments:
        raise SystemExit(f"Undefined segment values found: {undefined_segments[:20]}")
    if formula_errors:
        raise SystemExit(f"Formula errors found: {formula_errors[:20]}")
    if style_errors:
        raise SystemExit(f"Style errors found: {style_errors[:20]}")
    if row_count_errors:
        raise SystemExit(f"Row count errors found: {row_count_errors}")
    if forbidden_headers:
        raise SystemExit(f"Recommendation/action columns found: {forbidden_headers}")

    workbook.close()

    print(
        json.dumps(
            {
                "verified": str(args.workbook.resolve()),
                "sheets": expected_sheets,
                "control_totals": control_totals,
                "undefined_segments": 0,
                "formula_errors": 0,
                "font": f"{FONT_NAME} {FONT_SIZE} pt",
                "row_counts": expected_rows,
            },
            ensure_ascii=False,
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
