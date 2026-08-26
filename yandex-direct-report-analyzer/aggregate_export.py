#!/usr/bin/env python3
"""Нормализация и агрегация выгрузки контекстной рекламы в JSON для построения отчёта."""

from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from collections import defaultdict
from pathlib import Path


def norm(value: object) -> str:
    text = "" if value is None else str(value)
    text = text.replace("ё", "е").replace("Ё", "Е").strip().lower()
    text = re.sub(r"[\s\u00a0]+", " ", text)
    text = re.sub(r"[₽$€]", "", text)
    text = re.sub(r"[^0-9a-zа-я%]+", " ", text)
    return text.strip()


def clean(value: object) -> str:
    return "" if value is None else str(value).strip()


def number(value: object) -> float:
    if value is None or value == "":
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    raw = str(value).strip().replace("\u00a0", "").replace(" ", "")
    if raw in {"", "-", "—"}:
        return 0.0
    if raw.endswith("%"):
        raw = raw[:-1]
    if "," in raw and "." in raw:
        if raw.rfind(",") > raw.rfind("."):
            raw = raw.replace(".", "").replace(",", ".")
        else:
            raw = raw.replace(",", "")
    else:
        raw = raw.replace(",", ".")
    try:
        return float(raw)
    except ValueError:
        return 0.0


def rounded(value: float, digits: int = 6) -> float:
    return round(float(value), digits)


def is_undefined(value: object) -> bool:
    text = norm(value)
    return (
        text in {"", "неизвестно", "unknown", "undefined", "not set", "n a"}
        or text.startswith("не определ")
        or text.startswith("unknown")
    )


def identifier(value: object) -> str:
    text = clean(value)
    return "" if norm(text) == "0" or is_undefined(text) else text


FIELD_ALIASES = {
    "campaign": ["название кампании", "кампания", "campaign name"],
    "campaign_id": ["№ кампании", "номер кампании", "id кампании", "campaign id"],
    "group": ["название группы", "группа", "ad group", "group name"],
    "group_id": ["№ группы", "номер группы", "id группы", "group id"],
    "cost": ["расход", "расход руб", "затраты", "стоимость", "cost", "spend"],
    "clicks": ["клики", "clicks"],
    "impressions": ["показы", "impressions"],
    "revenue": ["выручка", "доход", "revenue", "conversion value", "ценность конверсий"],
    "orders": ["подтвержденные заказы", "оплаченные заказы", "заказы", "покупки", "orders", "purchases"],
}


SLICE_DEFS = [
    {"key": "condition_type", "label": "Тип условия показа", "sheet": "Тип условия показа", "aliases": ["тип условия показа"]},
    {"key": "keyword", "label": "Ключевая фраза", "sheet": "Ключевая фраза", "aliases": ["ключевая фраза", "фраза", "keyword"], "context": "group"},
    {"key": "brand_mention", "label": "Упоминание брендов в запросах", "sheet": "Упоминание бренда", "aliases": ["упоминания брендов в запросах", "упоминание брендов в запросах", "упоминание бренда", "брендовость запроса"]},
    {
        "key": "query_category",
        "label": "Категория запроса",
        "sheet": "Категория запроса",
        "aliases": ["категория запроса", "категория поискового запроса", "тематика запроса"],
        "context": "condition_source",
        "context_label": "Источник запроса",
        "context_in_identity": True,
    },
    {"key": "match_type", "label": "Тип соответствия", "sheet": "Тип соответствия", "aliases": ["тип соответствия", "тип соответствия фразы", "соответствие", "match type"], "context": "keyword"},
    {"key": "audience_condition", "label": "Условие подбора аудитории", "sheet": "Условие аудитории", "aliases": ["условие подбора аудитории", "аудитория"], "id_aliases": ["id условия подбора аудитории", "№ условия подбора аудитории"]},
    {"key": "target_adjustment", "label": "Корректировки для целевой аудитории", "sheet": "Корректировки ЦА", "aliases": ["название корректировки", "корректировки для целевой аудитории", "корректировка целевой аудитории"], "id_aliases": ["корректировки", "id корректировки", "№ корректировки"], "context": "audience"},
    {"key": "ad_number", "label": "№ объявления", "sheet": "№ объявления", "aliases": ["№ объявления", "номер объявления", "id объявления", "ad id"], "id_same": True, "context": "group"},
    {"key": "format", "label": "Формат", "sheet": "Формат", "aliases": ["формат", "формат объявления"], "context": "ad"},
    {"key": "region", "label": "Регион таргетинга", "sheet": "Регион таргетинга", "aliases": ["региона таргетинга", "регион таргетинга", "регион", "география"], "id_aliases": ["id региона таргетинга", "регион таргетинга id", "регион таргетинга"]},
    {"key": "device", "label": "Тип устройства", "sheet": "Тип устройства", "aliases": ["тип устройства", "устройство", "device type", "device"]},
    {"key": "age", "label": "Возраст", "sheet": "Возраст", "aliases": ["возраст", "age"]},
    {"key": "gender", "label": "Пол", "sheet": "Пол", "aliases": ["пол", "gender"]},
    {"key": "gender_age", "label": "Пол × возраст", "sheet": "Пол × возраст", "combine": ["gender", "age"]},
    {"key": "solvency", "label": "Уровень платёжеспособности", "sheet": "Платёжеспособность", "aliases": ["уровень платежеспособности", "уровень платёжеспособности", "платежеспособность", "платёжеспособность"]},
    {"key": "site_type", "label": "Тип площадки", "sheet": "Тип площадки", "aliases": ["тип площадки", "площадка"]},
    {"key": "placement", "label": "Вид размещения", "sheet": "Вид размещения", "aliases": ["вид размещения", "размещение"]},
    {"key": "group", "label": "Название группы", "sheet": "Название группы", "aliases": ["название группы", "группа", "ad group", "group name"], "id_aliases": ["№ группы", "номер группы", "id группы", "group id"]},
    {"key": "campaign_type", "label": "Тип кампании", "sheet": "Тип кампании", "aliases": ["тип кампании", "campaign type"]},
]


def header_candidates(headers: list[object]) -> list[str]:
    return [norm(item) for item in headers]


def find_all(headers: list[object], aliases: list[str]) -> list[int]:
    normalized = header_candidates(headers)
    wanted = [norm(alias) for alias in aliases]
    exact = [idx for idx, header in enumerate(normalized) if header in wanted]
    if exact:
        return exact
    return [idx for idx, header in enumerate(normalized) if any(alias and alias in header for alias in wanted)]


def find_one(headers: list[object], aliases: list[str], used: set[int] | None = None) -> int | None:
    used = used or set()
    for idx in find_all(headers, aliases):
        if idx not in used:
            return idx
    return None


def header_score(row: list[object]) -> int:
    headers = header_candidates(row)
    score = 0
    for aliases in (FIELD_ALIASES["cost"], FIELD_ALIASES["clicks"], FIELD_ALIASES["campaign"]):
        if any(norm(alias) in headers for alias in aliases):
            score += 2
    if any("конверс" in item for item in headers):
        score += 2
    score += min(3, sum(1 for spec in SLICE_DEFS if "aliases" in spec and any(norm(alias) in headers for alias in spec["aliases"])))
    return score


def csv_rows(path: Path):
    csv.field_size_limit(50_000_000)
    encoding = "utf-8-sig"
    try:
        with path.open("r", encoding=encoding, errors="strict") as probe:
            probe.read(65536)
    except UnicodeDecodeError:
        encoding = "cp1251"
    with path.open("r", encoding=encoding, newline="") as handle:
        sample = handle.read(65536)
        handle.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        yield from csv.reader(handle, dialect)


def xlsx_rows(path: Path, requested_sheet: str | None):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise SystemExit("XLSX input requires openpyxl in the loaded workspace dependencies") from exc
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook[requested_sheet] if requested_sheet else workbook[workbook.sheetnames[0]]
    for row in sheet.iter_rows(values_only=True):
        yield list(row)


def source_rows(path: Path, requested_sheet: str | None):
    suffix = path.suffix.lower()
    if suffix in {".csv", ".tsv", ".txt"}:
        return csv_rows(path)
    if suffix in {".xlsx", ".xlsm"}:
        return xlsx_rows(path, requested_sheet)
    raise SystemExit(f"Unsupported source format: {suffix}")


def metric_columns(headers: list[object]) -> dict:
    cost = find_one(headers, FIELD_ALIASES["cost"])
    clicks = find_one(headers, FIELD_ALIASES["clicks"])
    impressions = find_one(headers, FIELD_ALIASES["impressions"])
    revenue = find_one(headers, FIELD_ALIASES["revenue"])
    orders = find_one(headers, FIELD_ALIASES["orders"])
    normalized = header_candidates(headers)
    total_conversion_names = {"конверсии", "конверсии всего", "всего конверсий", "conversions", "total conversions"}
    total_conversion = next((idx for idx, item in enumerate(normalized) if item in total_conversion_names), None)
    if total_conversion is not None:
        conversions = [total_conversion]
    else:
        conversions = []
        for idx, item in enumerate(normalized):
            if "конверс" not in item and "conversion" not in item:
                continue
            if any(token in item for token in ("стоимость", "цена", "ценность", "доля", "коэффициент", "%", "rate", "cost", "value")):
                continue
            conversions.append(idx)
    if cost is None or clicks is None or not conversions:
        raise SystemExit("Could not map required Spend, Clicks, and Conversions columns")
    return {"cost": cost, "clicks": clicks, "impressions": impressions, "conversions": conversions, "revenue": revenue, "orders": orders}


def map_slices(headers: list[object]) -> tuple[list[dict], dict[str, int | None]]:
    base = {
        "campaign": find_one(headers, FIELD_ALIASES["campaign"]),
        "campaign_id": find_one(headers, FIELD_ALIASES["campaign_id"]),
        "group": find_one(headers, FIELD_ALIASES["group"]),
        "group_id": find_one(headers, FIELD_ALIASES["group_id"]),
    }
    mapped: list[dict] = []
    by_key: dict[str, int | None] = {}
    for spec in SLICE_DEFS:
        item = dict(spec)
        if "combine" in item:
            continue
        candidates = find_all(headers, item["aliases"])
        value_idx = candidates[0] if candidates else None
        if value_idx is None:
            by_key[item["key"]] = None
            continue
        item["value_idx"] = value_idx
        if item.get("id_same"):
            item["id_idx"] = value_idx
        elif item.get("id_aliases"):
            used = {value_idx}
            id_idx = find_one(headers, item["id_aliases"], used)
            if id_idx is None and len(candidates) > 1:
                id_idx = candidates[1]
            item["id_idx"] = id_idx
        else:
            item["id_idx"] = None
        mapped.append(item)
        by_key[item["key"]] = value_idx
    gender = by_key.get("gender")
    age = by_key.get("age")
    if gender is not None and age is not None:
        combo = next(spec for spec in SLICE_DEFS if spec["key"] == "gender_age")
        mapped.insert(next((i for i, item in enumerate(mapped) if item["key"] == "solvency"), len(mapped)), {**combo, "combine_idx": [gender, age]})
    return mapped, base


def value_at(row: list[object], idx: int | None) -> object:
    return "" if idx is None or idx >= len(row) else row[idx]


def metrics_for(row: list[object], columns: dict) -> dict[str, float]:
    return {
        "impressions": number(value_at(row, columns["impressions"])),
        "clicks": number(value_at(row, columns["clicks"])),
        "cost": number(value_at(row, columns["cost"])),
        "conversions": sum(number(value_at(row, idx)) for idx in columns["conversions"]),
        "revenue": number(value_at(row, columns["revenue"])),
        "orders": number(value_at(row, columns["orders"])),
    }


def add_metrics(target: dict[str, float], source: dict[str, float]) -> None:
    for key in ("impressions", "clicks", "cost", "conversions", "revenue", "orders"):
        target[key] += source[key]


def empty_metrics() -> dict[str, float]:
    return {key: 0.0 for key in ("impressions", "clicks", "cost", "conversions", "revenue", "orders")}


def is_inactive(metrics: dict[str, float]) -> bool:
    return all(metrics[key] == 0 for key in ("clicks", "cost", "conversions", "revenue", "orders"))


def make_context(spec: dict, row: list[object], base: dict[str, int | None], slice_map: dict[str, dict]) -> str:
    mode = spec.get("context")
    if mode == "group":
        name = clean(value_at(row, base["group"]))
        object_id = identifier(value_at(row, base["group_id"]))
    elif mode == "keyword":
        keyword = slice_map.get("keyword")
        name = clean(value_at(row, keyword["value_idx"])) if keyword else ""
        object_id = ""
    elif mode == "condition_source":
        condition = slice_map.get("condition_type")
        raw_name = clean(value_at(row, condition["value_idx"])) if condition else ""
        normalized = norm(raw_name)
        if "автотаргетинг" in normalized:
            name = "Автотаргетинг"
        elif "ключевая фраза" in normalized:
            name = "Ключевая фраза"
        else:
            name = raw_name
        object_id = ""
    elif mode == "audience":
        audience = slice_map.get("audience_condition")
        name = clean(value_at(row, audience["value_idx"])) if audience else ""
        object_id = identifier(value_at(row, audience.get("id_idx"))) if audience else ""
    elif mode == "ad":
        ad = slice_map.get("ad_number")
        name = clean(value_at(row, ad["value_idx"])) if ad else ""
        object_id = ""
    else:
        return ""
    if is_undefined(name):
        return ""
    return name + ((" | ID " + object_id) if object_id else "")


def get_segment(spec: dict, row: list[object]) -> tuple[str, str] | None:
    if "combine_idx" in spec:
        parts = [clean(value_at(row, idx)) for idx in spec["combine_idx"]]
        if any(is_undefined(part) for part in parts):
            return None
        return " × ".join(parts), ""
    value = clean(value_at(row, spec["value_idx"]))
    if is_undefined(value):
        return None
    return value, identifier(value_at(row, spec.get("id_idx")))


def finalize_metrics(source: dict[str, float]) -> dict[str, float]:
    return {
        key: rounded(source[key], 2 if key in {"cost", "revenue"} else 6)
        for key in ("impressions", "clicks", "cost", "conversions", "revenue", "orders")
    }


def new_acc(level: str, campaign: str, campaign_id: str, segment: str, segment_id: str, context: str) -> dict:
    return {"level": level, "campaign": campaign, "campaign_id": campaign_id, "segment": segment, "segment_id": segment_id, "context": context, **empty_metrics()}


def update_acc(acc: dict, metrics: dict[str, float], context: str) -> None:
    add_metrics(acc, metrics)
    if context and acc["context"] and acc["context"] != context:
        acc["context"] = "Несколько связанных объектов"
    elif context and not acc["context"]:
        acc["context"] = context


def finalize_acc(acc: dict) -> dict:
    result = {key: acc[key] for key in ("level", "campaign", "campaign_id", "segment", "segment_id", "context")}
    result.update(finalize_metrics(acc))
    return result


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True)
    parser.add_argument("--output-json", required=True)
    parser.add_argument("--sheet")
    args = parser.parse_args()

    source = Path(args.input).expanduser().resolve()
    output = Path(args.output_json).expanduser().resolve()
    if not source.is_file():
        raise SystemExit(f"Source file not found: {source}")

    iterator = iter(source_rows(source, args.sheet))
    prefix = []
    for _ in range(50):
        try:
            prefix.append(next(iterator))
        except StopIteration:
            break
    if not prefix:
        raise SystemExit("Source has no rows")
    header_index = max(range(len(prefix)), key=lambda idx: header_score(prefix[idx]))
    if header_score(prefix[header_index]) < 4:
        raise SystemExit("Could not identify a reliable header row")
    headers = [clean(item) for item in prefix[header_index]]
    metrics_columns = metric_columns(headers)
    slices, base = map_slices(headers)
    slice_map = {item["key"]: item for item in slices}

    campaign_acc: dict[str, dict] = {}
    campaign_seen: dict[str, tuple[str, str]] = {}
    inner_acc = {spec["key"]: {} for spec in slices}
    all_acc = {spec["key"]: {} for spec in slices}
    exclusions = {spec["key"]: {"rows": 0, **empty_metrics()} for spec in slices}
    detail_total = empty_metrics()
    source_total = None
    detail_rows = active_rows = inactive_rows = total_rows = 0
    conversions_gt_clicks_rows = 0
    conversions_gt_clicks_cost = 0.0

    data_rows = iter(prefix[header_index + 1 :] + list(iterator)) if source.stat().st_size < 50_000_000 else None
    if data_rows is None:
        def chained():
            yield from prefix[header_index + 1 :]
            yield from iterator
        data_rows = chained()

    for row in data_rows:
        if len(row) < len(headers):
            row = list(row) + [""] * (len(headers) - len(row))
        row_metrics = metrics_for(row, metrics_columns)
        first_labels = {norm(value_at(row, idx)) for idx in range(min(6, len(row)))}
        if first_labels.intersection({"итого", "всего", "total", "grand total"}):
            source_total = row_metrics
            total_rows += 1
            continue
        detail_rows += 1
        add_metrics(detail_total, row_metrics)
        campaign = clean(value_at(row, base["campaign"]))
        campaign_id = identifier(value_at(row, base["campaign_id"]))
        if is_undefined(campaign):
            campaign = "Кампания не указана"
        campaign_key = campaign_id or norm(campaign)
        campaign_seen[campaign_key] = (campaign, campaign_id)
        if row_metrics["conversions"] > row_metrics["clicks"] and row_metrics["clicks"] >= 0:
            conversions_gt_clicks_rows += 1
            conversions_gt_clicks_cost += row_metrics["cost"]
        if is_inactive(row_metrics):
            inactive_rows += 1
            continue
        active_rows += 1

        if campaign_key not in campaign_acc:
            campaign_acc[campaign_key] = new_acc("Кампания", campaign, campaign_id, campaign, campaign_id, "")
        update_acc(campaign_acc[campaign_key], row_metrics, "")

        for spec in slices:
            segment = get_segment(spec, row)
            if segment is None:
                exclusions[spec["key"]]["rows"] += 1
                add_metrics(exclusions[spec["key"]], row_metrics)
                continue
            segment_name, segment_id = segment
            context = make_context(spec, row, base, slice_map)
            identity = segment_id or norm(segment_name)
            if spec.get("context_in_identity"):
                if is_undefined(context):
                    exclusions[spec["key"]]["rows"] += 1
                    add_metrics(exclusions[spec["key"]], row_metrics)
                    continue
                identity = f"{norm(context)} | {identity}"
            inner_key = (campaign_key, identity)
            if inner_key not in inner_acc[spec["key"]]:
                inner_acc[spec["key"]][inner_key] = new_acc("Внутри РК", campaign, campaign_id, segment_name, segment_id, context)
            update_acc(inner_acc[spec["key"]][inner_key], row_metrics, context)
            # Группы объявлений существуют только внутри своих кампаний. Объединение одинаково
            # названных групп из разных кампаний в строку «Все РК» создаёт вводящую в заблуждение агрегацию.
            if spec["key"] != "group":
                if identity not in all_acc[spec["key"]]:
                    all_acc[spec["key"]][identity] = new_acc("Все РК", "Все РК", "", segment_name, segment_id, context)
                update_acc(all_acc[spec["key"]][identity], row_metrics, context)

    detail_total_final = finalize_metrics(detail_total)
    source_total_final = finalize_metrics(source_total or detail_total)
    for campaign_key, (campaign, campaign_id) in campaign_seen.items():
        if campaign_key not in campaign_acc:
            campaign_acc[campaign_key] = new_acc("Кампания", campaign, campaign_id, campaign, campaign_id, "")
    campaigns = [finalize_acc(item) for item in campaign_acc.values()]
    campaigns.sort(key=lambda item: (-item["cost"], item["campaign"]))
    output_slices = []
    for spec in slices:
        all_rows = [finalize_acc(item) for item in all_acc[spec["key"]].values()]
        inner_rows = [finalize_acc(item) for item in inner_acc[spec["key"]].values()]
        if not all_rows and not inner_rows:
            continue
        if spec.get("context_in_identity"):
            all_rows.sort(key=lambda item: (item["context"], -item["cost"], item["segment"]))
            inner_rows.sort(
                key=lambda item: (
                    item["campaign"],
                    item["context"],
                    -item["cost"],
                    item["segment"],
                )
            )
        else:
            all_rows.sort(key=lambda item: (-item["cost"], item["segment"]))
            inner_rows.sort(key=lambda item: (item["campaign"], -item["cost"], item["segment"]))
        excluded = exclusions[spec["key"]]
        output_slices.append({
            "key": spec["key"], "label": spec["label"], "sheet": spec["sheet"],
            "context_label": spec.get("context_label", "Контекст"),
            "rows": all_rows + inner_rows, "all_count": len(all_rows), "inner_count": len(inner_rows),
            "excluded_undefined": {"rows": excluded["rows"], **finalize_metrics(excluded)},
        })

    available_labels = [item["label"] for item in output_slices] + (["Название кампании"] if campaigns else [])
    requested_labels = [item["label"] for item in SLICE_DEFS] + ["Название кампании"]
    unavailable_labels = [label for label in requested_labels if label not in available_labels]
    revenue_available = metrics_columns["revenue"] is not None and (
        source_total_final["revenue"] > 0 or detail_total_final["revenue"] > 0
    )
    orders_available = revenue_available and metrics_columns["orders"] is not None and (
        source_total_final["orders"] > 0 or detail_total_final["orders"] > 0
    )
    unavailable_metrics = []
    if not revenue_available:
        unavailable_metrics.extend(["Выручка", "ДРР", "ROMI"])
    if not orders_available:
        unavailable_metrics.append("AOV")

    result = {
        "schema_version": 1,
        "source": str(source),
        "period": "Не указан в источнике",
        "headers": headers,
        "field_map": {
            "metrics": {key: ([headers[idx] for idx in value] if isinstance(value, list) else (headers[value] if value is not None else None)) for key, value in metrics_columns.items()},
            "dimensions": {spec["label"]: headers[spec["value_idx"]] if "value_idx" in spec else " + ".join(spec["combine"]) for spec in slices},
        },
        "metrics_available": {
            "revenue": revenue_available,
            "drr": revenue_available,
            "romi": revenue_available,
            "orders": orders_available,
            "aov": orders_available,
        },
        "available_dimensions": available_labels,
        "unavailable_dimensions": unavailable_labels,
        "unavailable_metrics": unavailable_metrics,
        "query_category_split": bool(
            slice_map.get("query_category") and slice_map.get("condition_type")
        ),
        "source_total": source_total_final,
        "detail_total": detail_total_final,
        "campaigns": campaigns,
        "slices": output_slices,
        "data_quality": {
            "header_row": header_index + 1,
            "detail_rows": detail_rows,
            "source_total_rows": total_rows,
            "active_rows": active_rows,
            "inactive_rows_excluded": inactive_rows,
            "rows_conversions_gt_clicks": conversions_gt_clicks_rows,
            "cost_conversions_gt_clicks": rounded(conversions_gt_clicks_cost, 2),
            "share_cost_on_rows_conversions_gt_clicks": rounded(conversions_gt_clicks_cost / detail_total["cost"], 8) if detail_total["cost"] else 0,
            "analysis_rows": len(campaigns) + sum(len(item["rows"]) for item in output_slices),
        },
    }
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(result, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(json.dumps({
        "source_total": source_total_final,
        "detail_total": detail_total_final,
        "campaigns": len(campaigns),
        "available_slices": available_labels,
        "unavailable_slices": unavailable_labels,
        "unavailable_metrics": unavailable_metrics,
        "data_quality": result["data_quality"],
    }, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    sys.exit(main())
