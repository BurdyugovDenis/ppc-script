#!/usr/bin/env python3
"""Построение Excel-отчёта с расчётами из нормализованного JSON с рекламными данными."""

from __future__ import annotations

import argparse
import json
import math
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter


FONT_NAME = "IBM Plex Sans"
FONT_SIZE = 10
WILSON_Z = 1.96
ERROR_THRESHOLD = 0.30
ZERO_CONVERSION_MULTIPLE = 3.0
SMALL_CR_CLICKS = 10
MAX_WRAPPED_LINES = 26
MAX_TEXT_COLUMN_WIDTH = 120

COLORS = {
    "dark": "173F35",
    "green": "2F7D62",
    "pale_green": "E8F4EE",
    "pale_blue": "EAF2FF",
    "amber": "FFF4CF",
    "red": "FDE7E7",
    "orange": "FFE8CC",
    "white": "FFFFFF",
    "text": "26352F",
    "line": "CFDBD5",
}

STATUS_STABLE = "Устойчиво"
STATUS_ERROR = "Гипотеза — погрешность > 30%"
STATUS_INSUFFICIENT = "Гипотеза — недостаточно данных"
STATUS_ZERO = "Гипотеза — нулевой расход"
STATUS_DATA = "Аномалия данных"
STATUS_X3_CAMPAIGN = "Аномалия — 0 конверсий при расходе >3× KPI CPL кампании"
STATUS_X3_ALL = "Аномалия — 0 конверсий при расходе >3× KPI CPL всех РК"
STATUS_X3_BOTH = "Аномалия — 0 конверсий при расходе >3× обоих KPI CPL"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--analysis", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    return parser.parse_args()


def safe_divide(numerator: float, denominator: float) -> float | None:
    return numerator / denominator if denominator else None


def order_count(row: dict, aov_from_conversions: bool) -> float:
    source = "conversions" if aov_from_conversions else "orders"
    return float(row.get(source) or 0)


def kpi(
    row: dict,
    drr_available: bool,
    romi_available: bool,
    aov_available: bool,
    aov_from_conversions: bool,
) -> dict:
    cost = float(row.get("cost") or 0)
    conversions = float(row.get("conversions") or 0)
    revenue = float(row.get("revenue") or 0)
    orders = order_count(row, aov_from_conversions)
    return {
        "cpl": safe_divide(cost, conversions),
        "drr": safe_divide(cost, revenue) if drr_available else None,
        "romi": safe_divide(revenue - cost, cost) if romi_available else None,
        "aov": safe_divide(revenue, orders) if aov_available else None,
    }


def safe_sheet_name(label: str, existing: set[str]) -> str:
    base = re.sub(r"[\\/*?:\[\]]", " ", label).strip()[:31] or "Срез"
    candidate = base
    suffix = 2
    while candidate in existing:
        tail = f" {suffix}"
        candidate = f"{base[:31 - len(tail)]}{tail}"
        suffix += 1
    existing.add(candidate)
    return candidate


def wilson_error(clicks: float, conversions: float) -> float | None:
    if clicks <= 0 or conversions <= 0 or conversions > clicks:
        return None
    cr = conversions / clicks
    numerator = WILSON_Z * math.sqrt(
        cr * (1 - cr) / clicks + WILSON_Z**2 / (4 * clicks**2)
    )
    denominator = cr + WILSON_Z**2 / (2 * clicks)
    return numerator / denominator if denominator else None


def score_row(
    row: dict,
    all_kpi: dict,
    campaign_kpis: dict[str, dict],
    drr_available: bool,
    romi_available: bool,
    aov_available: bool,
    aov_from_conversions: bool,
) -> dict:
    clicks = float(row.get("clicks") or 0)
    conversions = float(row.get("conversions") or 0)
    cost = float(row.get("cost") or 0)
    revenue = float(row.get("revenue") or 0)
    orders = order_count(row, aov_from_conversions)
    inner = row.get("level") == "Внутри РК"
    campaign_key = row.get("campaign_id") or row.get("campaign") or ""
    campaign_kpi = campaign_kpis.get(campaign_key, {}) if inner else {}

    cr = safe_divide(conversions, clicks)
    cpl = safe_divide(cost, conversions)
    drr = safe_divide(cost, revenue) if drr_available else None
    romi = safe_divide(revenue - cost, cost) if romi_available else None
    aov = safe_divide(revenue, orders) if aov_available else None
    error = wilson_error(clicks, conversions)
    multiple_campaign = (
        safe_divide(cost, campaign_kpi.get("cpl")) if campaign_kpi.get("cpl") else None
    )
    multiple_all = safe_divide(cost, all_kpi.get("cpl")) if all_kpi.get("cpl") else None

    if conversions > clicks:
        status = STATUS_DATA
    elif conversions == 0:
        if cost == 0:
            status = STATUS_ZERO
        else:
            above_campaign = inner and multiple_campaign is not None and multiple_campaign > ZERO_CONVERSION_MULTIPLE
            above_all = multiple_all is not None and multiple_all > ZERO_CONVERSION_MULTIPLE
            if above_campaign and above_all:
                status = STATUS_X3_BOTH
            elif above_campaign:
                status = STATUS_X3_CAMPAIGN
            elif above_all:
                status = STATUS_X3_ALL
            else:
                status = STATUS_INSUFFICIENT
    elif cost == 0:
        status = STATUS_ZERO
    elif clicks <= 0 or (cr == 1 and clicks < SMALL_CR_CLICKS):
        status = STATUS_INSUFFICIENT
    elif error is not None and error <= ERROR_THRESHOLD:
        status = STATUS_STABLE
    else:
        status = STATUS_ERROR

    stable = status == STATUS_STABLE
    correction_campaign = (
        campaign_kpi["cpl"] / cpl - 1
        if stable and inner and campaign_kpi.get("cpl") and cpl
        else None
    )
    correction_all = all_kpi["cpl"] / cpl - 1 if stable and all_kpi.get("cpl") and cpl else None
    drr_campaign = campaign_kpi.get("drr")
    drr_all = all_kpi.get("drr")
    romi_campaign = campaign_kpi.get("romi")
    romi_all = all_kpi.get("romi")
    aov_campaign = campaign_kpi.get("aov")
    aov_all = all_kpi.get("aov")

    correction_drr_campaign = (
        drr_campaign / drr - 1 if stable and inner and drr_campaign and drr else None
    )
    correction_drr_all = drr_all / drr - 1 if stable and drr_all and drr else None
    correction_romi_campaign = (
        (1 + romi) / (1 + romi_campaign) - 1
        if stable
        and inner
        and romi is not None
        and romi_campaign is not None
        and 1 + romi > 0
        and 1 + romi_campaign > 0
        else None
    )
    correction_romi_all = (
        (1 + romi) / (1 + romi_all) - 1
        if stable
        and romi is not None
        and romi_all is not None
        and 1 + romi > 0
        and 1 + romi_all > 0
        else None
    )
    aov_deviation_campaign = (
        aov / aov_campaign - 1 if inner and aov and aov_campaign else None
    )
    aov_deviation_all = aov / aov_all - 1 if aov and aov_all else None
    drr_deviation_campaign = (
        drr - drr_campaign
        if inner and drr is not None and drr_campaign is not None
        else None
    )
    drr_deviation_all = (
        drr - drr_all if drr is not None and drr_all is not None else None
    )
    romi_deviation_campaign = (
        romi - romi_campaign
        if inner and romi is not None and romi_campaign is not None
        else None
    )
    romi_deviation_all = (
        romi - romi_all if romi is not None and romi_all is not None else None
    )

    if stable:
        reason = ""
    elif status == STATUS_DATA:
        reason = "Конверсии превышают клики — проверьте состав целей и атрибуцию"
    elif conversions == 0:
        above_x3 = (
            (multiple_campaign is not None and multiple_campaign > ZERO_CONVERSION_MULTIPLE)
            or (multiple_all is not None and multiple_all > ZERO_CONVERSION_MULTIPLE)
        )
        reason = "0 конверсий; расход выше x3 KPI" if above_x3 else "Нет конверсий; порог x3 не превышен"
    elif cost == 0:
        reason = "Нулевой расход"
    elif cr == 1 and clicks < SMALL_CR_CLICKS:
        reason = "CR=100% на малой выборке"
    else:
        reason = "Погрешность выше 30%"

    return {
        "cr": cr,
        "cpl": cpl,
        "drr": drr,
        "romi": romi,
        "aov": aov,
        "orders": orders,
        "error": error,
        "multiple_campaign": multiple_campaign,
        "multiple_all": multiple_all,
        "correction_campaign": correction_campaign,
        "correction_all": correction_all,
        "correction_drr_campaign": correction_drr_campaign,
        "correction_drr_all": correction_drr_all,
        "correction_romi_campaign": correction_romi_campaign,
        "correction_romi_all": correction_romi_all,
        "drr_deviation_campaign": drr_deviation_campaign,
        "drr_deviation_all": drr_deviation_all,
        "romi_deviation_campaign": romi_deviation_campaign,
        "romi_deviation_all": romi_deviation_all,
        "aov_deviation_campaign": aov_deviation_campaign,
        "aov_deviation_all": aov_deviation_all,
        "campaign_kpi": campaign_kpi,
        "status": status,
        "reason": reason,
    }


def headers_for(
    revenue_available: bool,
    drr_available: bool,
    romi_available: bool,
    aov_available: bool,
    context_header: str = "Контекст",
) -> list[str]:
    headers = [
        "Уровень анализа",
        "Название кампании",
        "ID кампании",
        "Сегмент",
        "ID / номер сегмента",
        context_header,
        "Расход, ₽",
        "Доля расхода в скопе, %",
        "Клики",
        "Конверсии",
    ]
    if revenue_available:
        headers.append("Выручка, ₽")
    if aov_available:
        headers.append("Количество заказов")
    headers.extend(["CR", "CPL, ₽"])
    if drr_available:
        headers.append("ДРР")
    if romi_available:
        headers.append("ROMI")
    if aov_available:
        headers.append("AOV, ₽")
    headers.extend(
        [
            "Погрешность",
            "KPI CPL кампании, ₽",
            "KPI CPL всех РК, ₽",
            "Расход / KPI CPL кампании, ×",
            "Расход / KPI CPL всех РК, ×",
            "Расчёт к РК",
            "Расчёт к всем РК",
        ]
    )
    if drr_available:
        headers.extend(
            [
                "KPI ДРР кампании",
                "KPI ДРР всех РК",
                "Отклонение ДРР к РК, п.п.",
                "Отклонение ДРР к всем РК, п.п.",
                "Расчёт к РК по ДРР",
                "Расчёт к всем РК по ДРР",
            ]
        )
    if romi_available:
        headers.extend(
            [
                "KPI ROMI кампании",
                "KPI ROMI всех РК",
                "Отклонение ROMI к РК, п.п.",
                "Отклонение ROMI к всем РК, п.п.",
                "Расчёт к РК по ROMI",
                "Расчёт к всем РК по ROMI",
            ]
        )
    if aov_available:
        headers.extend(
            [
                "KPI AOV кампании, ₽",
                "KPI AOV всех РК, ₽",
                "Отклонение AOV к РК",
                "Отклонение AOV к всем РК",
            ]
        )
    headers.extend(["Статус данных", "Причина отсутствия расчёта"])
    return headers


def row_values(
    row: dict,
    calculation: dict,
    headers: list[str],
    all_kpi: dict,
    total_cost: float,
    campaign_costs: dict[str, float],
    revenue_available: bool,
    drr_available: bool,
    romi_available: bool,
    aov_available: bool,
    campaign_sheet: bool,
) -> list:
    inner = row.get("level") == "Внутри РК"
    campaign_key = row.get("campaign_id") or row.get("campaign") or ""
    scope_cost = campaign_costs.get(campaign_key, 0) if inner else total_cost
    campaign_kpi = calculation["campaign_kpi"] if inner and not campaign_sheet else {}
    values = {
        "Уровень анализа": row.get("level"),
        "Название кампании": row.get("campaign"),
        "ID кампании": row.get("campaign_id") or None,
        "Сегмент": row.get("segment"),
        "ID / номер сегмента": row.get("segment_id") or None,
        "Контекст": row.get("context") or None,
        "Источник запроса": row.get("context") or None,
        "Расход, ₽": float(row.get("cost") or 0),
        "Доля расхода в скопе, %": safe_divide(float(row.get("cost") or 0), scope_cost),
        "Клики": float(row.get("clicks") or 0),
        "Конверсии": float(row.get("conversions") or 0),
        "CR": calculation["cr"],
        "CPL, ₽": calculation["cpl"],
        "Погрешность": calculation["error"],
        "KPI CPL кампании, ₽": campaign_kpi.get("cpl"),
        "KPI CPL всех РК, ₽": all_kpi.get("cpl"),
        "Расход / KPI CPL кампании, ×": calculation["multiple_campaign"],
        "Расход / KPI CPL всех РК, ×": calculation["multiple_all"],
        "Расчёт к РК": calculation["correction_campaign"],
        "Расчёт к всем РК": calculation["correction_all"],
        "Статус данных": calculation["status"],
        "Причина отсутствия расчёта": calculation["reason"],
    }
    if revenue_available:
        values["Выручка, ₽"] = float(row.get("revenue") or 0)
    if drr_available:
        values.update(
            {
                "ДРР": calculation["drr"],
                "KPI ДРР кампании": campaign_kpi.get("drr"),
                "KPI ДРР всех РК": all_kpi.get("drr"),
                "Отклонение ДРР к РК, п.п.": calculation["drr_deviation_campaign"],
                "Отклонение ДРР к всем РК, п.п.": calculation["drr_deviation_all"],
                "Расчёт к РК по ДРР": calculation["correction_drr_campaign"],
                "Расчёт к всем РК по ДРР": calculation["correction_drr_all"],
            }
        )
    if romi_available:
        values.update(
            {
                "ROMI": calculation["romi"],
                "KPI ROMI кампании": campaign_kpi.get("romi"),
                "KPI ROMI всех РК": all_kpi.get("romi"),
                "Отклонение ROMI к РК, п.п.": calculation["romi_deviation_campaign"],
                "Отклонение ROMI к всем РК, п.п.": calculation["romi_deviation_all"],
                "Расчёт к РК по ROMI": calculation["correction_romi_campaign"],
                "Расчёт к всем РК по ROMI": calculation["correction_romi_all"],
            }
        )
    if aov_available:
        values.update(
            {
                "Количество заказов": calculation["orders"],
                "AOV, ₽": calculation["aov"],
                "KPI AOV кампании, ₽": campaign_kpi.get("aov"),
                "KPI AOV всех РК, ₽": all_kpi.get("aov"),
                "Отклонение AOV к РК": calculation["aov_deviation_campaign"],
                "Отклонение AOV к всем РК": calculation["aov_deviation_all"],
            }
        )
    return [values.get(header) for header in headers]


def style_header(row_cells) -> None:
    fill = PatternFill("solid", fgColor=COLORS["dark"])
    font = Font(name=FONT_NAME, size=FONT_SIZE, bold=True, color=COLORS["white"])
    side = Side(style="medium", color=COLORS["dark"])
    for cell in row_cells:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(left=side, right=side, top=side, bottom=side)


def apply_base_style(ws) -> None:
    thin = Side(style="thin", color=COLORS["line"])
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            cell.font = Font(
                name=FONT_NAME,
                size=FONT_SIZE,
                bold=cell.font.bold,
                italic=cell.font.italic,
                color=cell.font.color,
            )
            cell.alignment = Alignment(
                horizontal=cell.alignment.horizontal,
                vertical="center",
                wrap_text=cell.alignment.wrap_text,
            )
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws.sheet_view.showGridLines = False


def apply_analysis_grid(ws, headers: list[str]) -> None:
    """Нанести непрерывную видимую сетку и единое оформление текста на всю таблицу."""
    thin = Side(style="thin", color="A8B4AE")
    outer = Side(style="medium", color=COLORS["dark"])
    body_font = Font(name=FONT_NAME, size=FONT_SIZE, color=COLORS["text"])
    plain_alignment = Alignment(vertical="center")
    wrapped_alignment = Alignment(vertical="center", wrap_text=True)
    wrapped_headers = {
        "Название кампании",
        "Сегмент",
        "Контекст",
        "Источник запроса",
        "Статус данных",
        "Причина отсутствия расчёта",
    }
    wrapped_columns = {
        index + 1 for index, header in enumerate(headers) if header in wrapped_headers
    }
    max_row = ws.max_row
    max_column = ws.max_column
    for row_number, row in enumerate(
        ws.iter_rows(min_row=1, max_row=max_row, min_col=1, max_col=max_column),
        start=1,
    ):
        for column_number, cell in enumerate(row, start=1):
            cell.font = body_font
            cell.alignment = (
                wrapped_alignment if column_number in wrapped_columns else plain_alignment
            )
            cell.border = Border(
                left=outer if column_number == 1 else thin,
                right=outer if column_number == max_column else thin,
                top=outer if row_number == 1 else thin,
                bottom=outer if row_number == max_row else thin,
            )
    ws.sheet_view.showGridLines = False


def fit_row_heights(
    ws,
    wrapped_columns: set[int],
    header_height: float = 36,
) -> None:
    """Рассчитать высоту строк Excel после определения ширины столбцов и переноса текста."""
    for row_number in range(1, ws.max_row + 1):
        if row_number == 1:
            ws.row_dimensions[row_number].height = header_height
            continue
        max_lines = 1
        for column_index in wrapped_columns:
            cell = ws.cell(row_number, column_index)
            if cell.value is None:
                continue
            width = float(ws.column_dimensions[get_column_letter(column_index)].width or 10)
            # IBM Plex Sans и перенос по словам занимают немного больше места, чем показывает
            # простой подсчёт символов. Запас по ширине не даёт обрезаться краям текста.
            usable_width = max((width - 1) * 0.90, 1)
            parts = str(cell.value).splitlines() or [""]
            line_count = sum(
                max(1, math.ceil(len(part) / usable_width))
                for part in parts
            )
            max_lines = max(max_lines, line_count)
        # Excel поддерживает высоту строки до 409,5 пт. Используем весь доступный предел,
        # чтобы длинные фразы и условия аудиторий не обрезались после шести строк.
        ws.row_dimensions[row_number].height = min(
            409.5, max(18, 15 * max_lines + 3)
        )


def autofit(ws, headers: list[str] | None = None) -> None:
    text_headers = {
        "Название кампании",
        "ID кампании",
        "Сегмент",
        "Контекст",
        "Источник запроса",
        "Статус данных",
        "Причина отсутствия расчёта",
    }
    preferred_widths = {
        "Название кампании": 36,
        "ID кампании": 18,
        "Сегмент": 46,
        "Контекст": 36,
        "Источник запроса": 24,
        "Статус данных": 34,
        "Причина отсутствия расчёта": 40,
    }
    header_by_column = {index + 1: value for index, value in enumerate(headers or [])}
    for column_cells in ws.iter_cols(1, ws.max_column, 1, ws.max_row):
        column_index = column_cells[0].column
        header = header_by_column.get(column_index, "")
        if header in text_headers:
            width = max(preferred_widths.get(header, 32), len(str(header)) + 2)
            for cell in column_cells:
                if cell.value is None:
                    continue
                parts = str(cell.value).splitlines() or [""]
                # Увеличиваем ширину только тогда, когда текущее значение потребовало бы высоту
                # больше максимума Excel. Благодаря этому обычные листы остаются компактными.
                while width < MAX_TEXT_COLUMN_WIDTH:
                    usable_width = max((width - 1) * 0.90, 1)
                    lines = sum(
                        max(1, math.ceil(len(part) / usable_width))
                        for part in parts
                    )
                    if lines <= MAX_WRAPPED_LINES:
                        break
                    width = min(MAX_TEXT_COLUMN_WIDTH, width + max(2, width * 0.08))
            final_width = min(width, MAX_TEXT_COLUMN_WIDTH)
        else:
            maximum = 0
            for cell in column_cells[:250]:
                value = "" if cell.value is None else str(cell.value)
                maximum = max(
                    maximum,
                    max((len(part) for part in value.splitlines()), default=0),
                )
            final_width = min(max(maximum + 2, 10), 24)
        ws.column_dimensions[get_column_letter(column_index)].width = final_width
    wrapped_columns = {
        index
        for index, header in header_by_column.items()
        if header in text_headers or header == "Статус данных"
    }
    fit_row_heights(ws, wrapped_columns)


def format_analysis_sheet(ws, headers: list[str], table_name: str) -> None:
    apply_analysis_grid(ws, headers)
    style_header(ws[1])
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = ws.dimensions
    if ws.max_row >= 2:
        table = Table(displayName=table_name, ref=ws.dimensions)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleLight9",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=False,
            showColumnStripes=False,
        )
        ws.add_table(table)
    money = {
        "Расход, ₽",
        "Выручка, ₽",
        "CPL, ₽",
        "AOV, ₽",
        "KPI CPL кампании, ₽",
        "KPI CPL всех РК, ₽",
        "KPI AOV кампании, ₽",
        "KPI AOV всех РК, ₽",
    }
    percentages = {
        "Доля расхода в скопе, %",
        "CR",
        "Погрешность",
        "ДРР",
        "ROMI",
        "KPI ДРР кампании",
        "KPI ДРР всех РК",
        "KPI ROMI кампании",
        "KPI ROMI всех РК",
    }
    adjustments = {header for header in headers if header.startswith("Расчёт") or header.startswith("Отклонение")}
    multiples = {header for header in headers if "KPI CPL" in header and header.endswith("×")}
    status_col = headers.index("Статус данных") + 1
    for index, header in enumerate(headers, 1):
        number_format = None
        if header in money:
            number_format = '#,##0.00 "₽"'
        elif header in percentages:
            number_format = "0.00%"
        elif header in adjustments:
            number_format = "+0.0%;-0.0%;0.0%"
        elif header in multiples:
            number_format = '0.00"×"'
        elif header in {"Клики", "Конверсии", "Количество заказов"}:
            number_format = "#,##0"
        if number_format:
            for (target,) in ws.iter_rows(
                min_row=2,
                max_row=ws.max_row,
                min_col=index,
                max_col=index,
            ):
                target.number_format = number_format
    for row in range(2, ws.max_row + 1):
        status_cell = ws.cell(row, status_col)
        status = str(status_cell.value or "")
        if status == STATUS_STABLE:
            status_cell.fill = PatternFill("solid", fgColor="DCFCE7")
            status_cell.font = Font(name=FONT_NAME, size=FONT_SIZE, bold=True, color="166534")
        elif status.startswith("Аномалия"):
            status_cell.fill = PatternFill("solid", fgColor=COLORS["red"])
            status_cell.font = Font(name=FONT_NAME, size=FONT_SIZE, bold=True, color="991B1B")
        else:
            status_cell.fill = PatternFill("solid", fgColor=COLORS["amber"])
            status_cell.font = Font(name=FONT_NAME, size=FONT_SIZE, color="7C4A03")
    for header in adjustments:
        column = headers.index(header) + 1
        letter = get_column_letter(column)
        ws.conditional_formatting.add(
            f"{letter}2:{letter}{ws.max_row}",
            CellIsRule(operator="greaterThan", formula=["0"], fill=PatternFill("solid", fgColor="DCFCE7")),
        )
        ws.conditional_formatting.add(
            f"{letter}2:{letter}{ws.max_row}",
            CellIsRule(operator="lessThan", formula=["0"], fill=PatternFill("solid", fgColor="FEE2E2")),
        )
    style_header(ws[1])
    autofit(ws, headers)


def write_summary(
    ws,
    data: dict,
    all_kpi: dict,
    counts: dict,
    sheet_rows: list[tuple],
    aov_from_conversions: bool,
) -> None:
    ws.merge_cells("A1:F1")
    ws["A1"] = "Расчёт показателей контекстной рекламы"
    ws.merge_cells("A2:F2")
    ws["A2"] = f"Период: {data.get('period') or 'Не указан'}"
    ws.merge_cells("A3:F3")
    ws["A3"] = f"Источник: {data.get('source') or ''}"
    metrics = data.get("metrics_available", {})
    summary_rows = [
        ("Показатель", "Значение"),
        ("Расход, ₽", data["source_total"].get("cost")),
        ("Клики", data["source_total"].get("clicks")),
        ("Конверсии", data["source_total"].get("conversions")),
        ("CR", safe_divide(data["source_total"].get("conversions", 0), data["source_total"].get("clicks", 0))),
        ("CPL, ₽", all_kpi.get("cpl")),
    ]
    if metrics.get("revenue"):
        summary_rows.append(("Выручка, ₽", data["source_total"].get("revenue")))
    if metrics.get("drr"):
        summary_rows.append(("ДРР", all_kpi.get("drr")))
    if metrics.get("romi"):
        summary_rows.append(("ROMI", all_kpi.get("romi")))
    if metrics.get("aov"):
        summary_rows.extend([
            (
                "Конверсии как заказы для AOV"
                if aov_from_conversions
                else "Количество заказов",
                order_count(data["source_total"], aov_from_conversions),
            ),
            ("AOV, ₽", all_kpi.get("aov")),
        ])
    for row in summary_rows:
        ws.append(row)
    start = ws.max_row + 2
    ws.cell(start, 1, "Статус расчётов")
    ws.cell(start, 2, "Количество")
    for label, value in [
        ("Кампаний", len(data.get("campaigns", []))),
        ("Аналитических строк", data.get("data_quality", {}).get("analysis_rows", 0)),
        ("Устойчивых строк", counts["stable"]),
        ("Гипотез", counts["hypotheses"]),
        ("Аномалий", counts["anomalies"]),
        ("x3-аномалий", counts["x3"]),
    ]:
        ws.append((label, value))
    start = ws.max_row + 2
    ws.cell(start, 1, "Лист")
    ws.cell(start, 2, "Срез")
    ws.cell(start, 3, "Строк")
    ws.cell(start, 4, "Исключено неопределённых строк")
    for item in sheet_rows:
        ws.append(item)
    ws.freeze_panes = "A4"
    ws.sheet_view.showGridLines = False
    for row_number in (4, start):
        style_header(ws[row_number])
    title_fill = PatternFill("solid", fgColor=COLORS["dark"])
    ws["A1"].fill = title_fill
    ws["A1"].font = Font(name=FONT_NAME, size=FONT_SIZE, bold=True, color=COLORS["white"])
    ws["A2"].fill = PatternFill("solid", fgColor=COLORS["pale_green"])
    apply_base_style(ws)
    ws["A1"].fill = title_fill
    ws["A1"].font = Font(name=FONT_NAME, size=FONT_SIZE, bold=True, color=COLORS["white"])
    status_header_row = next(
        row
        for row in range(1, ws.max_row + 1)
        if ws.cell(row, 1).value == "Статус расчётов"
    )
    for row in range(5, status_header_row):
        label = ws.cell(row, 1).value
        if label in {"Расход, ₽", "CPL, ₽", "Выручка, ₽", "AOV, ₽"}:
            ws.cell(row, 2).number_format = '#,##0.00 "₽"'
        elif label in {"CR", "ДРР", "ROMI"}:
            ws.cell(row, 2).number_format = "0.00%"
        elif label in {
            "Клики",
            "Конверсии",
            "Количество заказов",
            "Конверсии как заказы для AOV",
        }:
            ws.cell(row, 2).number_format = "#,##0"
    sheet_header_row = next(
        row
        for row in range(1, ws.max_row + 1)
        if ws.cell(row, 1).value == "Лист"
    )
    for row in range(status_header_row + 1, sheet_header_row - 1):
        ws.cell(row, 2).number_format = "#,##0"
    for row in range(sheet_header_row + 1, ws.max_row + 1):
        ws.cell(row, 3).number_format = "#,##0"
        ws.cell(row, 4).number_format = "#,##0"
    autofit(ws)
    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 42
    ws.column_dimensions["D"].width = 34
    ws["A2"].alignment = Alignment(vertical="center", wrap_text=True)
    ws["A3"].alignment = Alignment(vertical="center", wrap_text=True)
    fit_row_heights(ws, {1, 2, 4}, header_height=30)


def write_method(ws, data: dict, all_kpi: dict) -> None:
    metrics = data.get("metrics_available", {})
    rows: list[tuple[str, object]] = [
        ("Параметр", "Значение"),
        ("Порог статистической погрешности", ERROR_THRESHOLD),
        ("Порог расхода без конверсий, × KPI CPL", ZERO_CONVERSION_MULTIPLE),
        ("Малая выборка при CR=100%, кликов", SMALL_CR_CLICKS),
        ("z для интервала Уилсона, 95%", WILSON_Z),
        ("Источник", data.get("source")),
        ("Период", data.get("period")),
        ("Общий расход, ₽", data["source_total"].get("cost")),
        ("Общие клики", data["source_total"].get("clicks")),
        ("Общие конверсии", data["source_total"].get("conversions")),
        ("KPI CPL всех РК, ₽", all_kpi.get("cpl")),
    ]
    if metrics.get("revenue"):
        rows.append(("Общая выручка, ₽", data["source_total"].get("revenue")))
    if metrics.get("drr"):
        rows.append(("KPI ДРР всех РК", all_kpi.get("drr")))
    if metrics.get("romi"):
        rows.append(("KPI ROMI всех РК", all_kpi.get("romi")))
    if metrics.get("aov"):
        rows.append(("KPI AOV всех РК, ₽", all_kpi.get("aov")))
    rows.extend([
        ("Формула", "Определение"),
        ("CR", "Конверсии / Клики"),
        ("CPL", "Расход / Конверсии"),
        ("Погрешность Wilson 95%", "z×√(CR×(1−CR)/Клики+z²/(4×Клики²))/(CR+z²/(2×Клики))"),
        ("Расчёт к РК", "KPI кампании / показатель сегмента − 1"),
        ("Расчёт к всем РК", "KPI всех РК / показатель сегмента − 1"),
    ])
    if metrics.get("drr"):
        rows.extend([
            ("ДРР", "Расход / Выручка"),
            ("Отклонение ДРР", "ДРР сегмента − KPI ДРР; в процентных пунктах"),
        ])
    if metrics.get("romi"):
        rows.extend([
            ("ROMI", "(Выручка − Расход) / Расход"),
            ("Отклонение ROMI", "ROMI сегмента − KPI ROMI; в процентных пунктах"),
        ])
    if metrics.get("aov"):
        rows.extend([
            ("AOV", data.get("aov_basis") or "Выручка / подтверждённые заказы"),
            ("Отклонение AOV", "AOV сегмента / KPI AOV − 1"),
        ])
    rows.extend([
        ("Правило устойчивости", "Погрешность ≤30%; CR=100% требует не менее 10 кликов"),
        ("Неопределённые сегменты", "Исключены из таблиц, но сохранены в KPI"),
        ("Недоступные срезы", ", ".join(data.get("unavailable_dimensions", [])) or "Нет"),
        ("Недоступные метрики", ", ".join(data.get("unavailable_metrics", [])) or "Нет"),
        ("Отключённые метрики", ", ".join(data.get("disabled_metrics", [])) or "Нет"),
        ("Детальных строк", data.get("data_quality", {}).get("detail_rows")),
        ("Неактивных строк исключено", data.get("data_quality", {}).get("inactive_rows_excluded")),
        ("Строк с конверсиями > кликов", data.get("data_quality", {}).get("rows_conversions_gt_clicks")),
        ("Расхождение итогов, ₽", data["detail_total"].get("cost", 0) - data["source_total"].get("cost", 0)),
        ("Исходные столбцы", " · ".join(data.get("headers", []))),
        ("База AOV", data.get("aov_basis") or "Недоступна"),
        (
            "Разделение категорий запросов",
            (
                "Категории сгруппированы отдельно по автотаргетингу, "
                "ключевым фразам и другим исходным типам условия показа"
                if data.get("query_category_split")
                else "Недоступно: нет типа условия показа или категории запроса"
            ),
        ),
    ])
    for row in rows:
        ws.append(row)
    formula_header_row = next(
        row for row in range(1, ws.max_row + 1)
        if ws.cell(row, 1).value == "Формула"
    )
    style_header(ws[1])
    style_header(ws[formula_header_row])
    for row in range(2, ws.max_row + 1):
        label = ws.cell(row, 1).value
        if label == "Порог статистической погрешности":
            ws.cell(row, 2).number_format = "0%"
        elif label in {"Общий расход, ₽", "Общая выручка, ₽", "KPI CPL всех РК, ₽", "KPI AOV всех РК, ₽", "Расхождение итогов, ₽"}:
            ws.cell(row, 2).number_format = '#,##0.00 "₽"'
        elif label in {"KPI ДРР всех РК", "KPI ROMI всех РК"}:
            ws.cell(row, 2).number_format = "0.00%"
        elif label in {"Малая выборка при CR=100%, кликов", "Общие клики", "Общие конверсии", "Детальных строк", "Неактивных строк исключено", "Строк с конверсиями > кликов"}:
            ws.cell(row, 2).number_format = "#,##0"
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False
    apply_base_style(ws)
    style_header(ws[1])
    style_header(ws[formula_header_row])
    autofit(ws)
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 86
    for row in range(1, ws.max_row + 1):
        ws.cell(row, 1).alignment = Alignment(vertical="center", wrap_text=True)
        ws.cell(row, 2).alignment = Alignment(vertical="center", wrap_text=True)
    fit_row_heights(ws, {1, 2}, header_height=30)


def main() -> None:
    args = parse_args()
    data = json.loads(args.analysis.read_text(encoding="utf-8"))
    for slice_data in data.get("slices", []):
        if slice_data.get("key") == "group":
            slice_data["rows"] = [row for row in slice_data.get("rows", []) if row.get("level") != "Все РК"]

    revenue_available = bool(data.get("metrics_available", {}).get("revenue"))
    drr_available = bool(data.get("metrics_available", {}).get("drr"))
    romi_available = bool(data.get("metrics_available", {}).get("romi"))
    aov_from_conversions = bool(data.get("aov_from_conversions"))
    aov_available = bool(
        revenue_available
        and data.get("metrics_available", {}).get("aov")
        and (
            data.get("metrics_available", {}).get("orders")
            or aov_from_conversions
        )
    )
    all_kpi = kpi(
        data["source_total"],
        drr_available,
        romi_available,
        aov_available,
        aov_from_conversions,
    )
    campaign_kpis = {
        row.get("campaign_id") or row.get("campaign"): kpi(
            row,
            drr_available,
            romi_available,
            aov_available,
            aov_from_conversions,
        )
        for row in data.get("campaigns", [])
    }
    campaign_costs = {
        row.get("campaign_id") or row.get("campaign"): float(row.get("cost") or 0)
        for row in data.get("campaigns", [])
    }
    total_cost = float(data["source_total"].get("cost") or 0)
    workbook = Workbook()
    workbook.remove(workbook.active)
    summary = workbook.create_sheet("Резюме")
    method = workbook.create_sheet("Методика")
    campaign_sheet = workbook.create_sheet("Название кампании")
    existing = {"Резюме", "Методика", "Название кампании"}
    slice_names = {
        slice_data["key"]: safe_sheet_name(slice_data.get("sheet") or slice_data["label"], existing)
        for slice_data in data.get("slices", [])
    }
    for slice_data in data.get("slices", []):
        workbook.create_sheet(slice_names[slice_data["key"]])

    counts = {"stable": 0, "hypotheses": 0, "anomalies": 0, "x3": 0}

    def write_rows(
        ws,
        rows: list[dict],
        table_index: int,
        campaign_sheet_flag: bool = False,
        context_header: str = "Контекст",
    ) -> None:
        headers = headers_for(
            revenue_available,
            drr_available,
            romi_available,
            aov_available,
            context_header,
        )
        ws.append(headers)
        for row in rows:
            calculation = score_row(
                row,
                all_kpi,
                campaign_kpis,
                drr_available,
                romi_available,
                aov_available,
                aov_from_conversions,
            )
            status = calculation["status"]
            if status == STATUS_STABLE:
                counts["stable"] += 1
            elif status.startswith("Аномалия"):
                counts["anomalies"] += 1
                if status in {STATUS_X3_CAMPAIGN, STATUS_X3_ALL, STATUS_X3_BOTH}:
                    counts["x3"] += 1
            else:
                counts["hypotheses"] += 1
            ws.append(
                row_values(
                    row,
                    calculation,
                    headers,
                    all_kpi,
                    total_cost,
                    campaign_costs,
                    revenue_available,
                    drr_available,
                    romi_available,
                    aov_available,
                    campaign_sheet_flag,
                )
            )
        format_analysis_sheet(ws, headers, f"AnalysisTable{table_index:03d}")

    write_rows(campaign_sheet, data.get("campaigns", []), 1, True)
    sheet_rows = [("Название кампании", "Название кампании", len(data.get("campaigns", [])), 0)]
    for index, slice_data in enumerate(data.get("slices", []), 2):
        sheet_name = slice_names[slice_data["key"]]
        rows = slice_data.get("rows", [])
        write_rows(
            workbook[sheet_name],
            rows,
            index,
            context_header=slice_data.get("context_label") or "Контекст",
        )
        sheet_rows.append(
            (
                sheet_name,
                slice_data["label"],
                len(rows),
                slice_data.get("excluded_undefined", {}).get("rows", 0),
            )
        )

    write_summary(
        summary,
        data,
        all_kpi,
        counts,
        sheet_rows,
        aov_from_conversions,
    )
    write_method(method, data, all_kpi)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(args.output)
    print(
        json.dumps(
            {
                "output": str(args.output.resolve()),
                "sheets": workbook.sheetnames,
                "analysis_rows": data.get("data_quality", {}).get("analysis_rows", 0),
                **counts,
            },
            ensure_ascii=False,
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
