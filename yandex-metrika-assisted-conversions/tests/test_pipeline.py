from __future__ import annotations

import csv
import json
import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from yandex_attribution.config import Settings, parse_goal_ids
from yandex_attribution.cli import (
    DIRECT_REPORT_CACHE_VERSION,
    METRIKA_LOG_CACHE_VERSION,
    direct_cache_matches,
    env_arguments,
    make_parser,
    metrika_cache_matches,
    output_filename,
)
from yandex_attribution.pipeline import (
    channel_assists_sql,
    conversion_touches_sql,
    association_window,
    build,
    build_conversion_touches,
    connect,
    create_performance_tables,
    create_schema,
    load_metrika_visits,
)


class PipelineTest(unittest.TestCase):
    @patch("yandex_attribution.cli.load_dotenv")
    def test_pycharm_env_uses_same_parser_for_all_modes(self, _load_env) -> None:
        for mode in ("run", "build", "extract", "goals"):
            with self.subTest(mode=mode), patch.dict("os.environ", {
                "RUN_MODE": f" {mode.upper()} ",
                "DATE_FROM": " 2026-08-01 ",
                "DATE_TO": "2026-08-31",
                "DATA_DIR": " report data ",
            }, clear=True):
                args = make_parser().parse_args(env_arguments())
                self.assertEqual(args.command, mode)
                self.assertEqual(args.data_dir, "report data")
                if mode != "goals":
                    self.assertEqual(args.date_from, date(2026, 8, 1))
                    self.assertEqual(args.date_to, date(2026, 8, 31))
                if mode in ("run", "extract"):
                    self.assertFalse(args.force)

    @patch("yandex_attribution.cli.load_dotenv")
    def test_pycharm_goals_does_not_require_dates(self, _load_env) -> None:
        with patch.dict("os.environ", {"RUN_MODE": "goals"}, clear=True):
            self.assertEqual(env_arguments(), ["--data-dir", "data", "goals"])

    @patch("yandex_attribution.cli.load_dotenv")
    def test_pycharm_rejects_missing_dates_and_invalid_mode(self, _load_env) -> None:
        for values, message in (
            ({}, "DATE_FROM"),
            ({"DATE_FROM": "2026-08-01"}, "DATE_TO"),
            ({"RUN_MODE": "unknown"}, "RUN_MODE"),
        ):
            with self.subTest(values=values), patch.dict("os.environ", values, clear=True):
                with self.assertRaisesRegex(ValueError, message):
                    env_arguments()

    def test_association_window_covers_per_conversion_history(self) -> None:
        self.assertEqual(
            association_window(date(2026, 8, 1), date(2026, 8, 31), 90),
            (date(2026, 5, 3), date(2026, 8, 31)),
        )

    def test_goal_ids_are_comma_separated_and_deduplicated(self) -> None:
        self.assertEqual(parse_goal_ids("99, 100,99"), (99, 100))

    def test_output_filename_contains_period_and_client_login(self) -> None:
        self.assertEqual(
            output_filename(
                date(2026, 8, 1), date(2026, 8, 31), "client-login"
            ),
            "2026-08-01_2026-08-31_client-login.xlsx",
        )
        self.assertEqual(
            output_filename(
                date(2026, 8, 1), date(2026, 8, 31), " agency/client: one "
            ),
            "2026-08-01_2026-08-31_agency_client_one.xlsx",
        )

    def test_all_settings_can_be_loaded_from_one_env_mapping(self) -> None:
        settings = Settings.from_env(
            {
                "YANDEX_METRIKA_COUNTER_IDS": "123,456,123",
                "YANDEX_DIRECT_CLIENT_LOGIN": "client-login",
                "YANDEX_GOAL_IDS": "99,100",
                "LOOKBACK_DAYS": "45",
                "DIRECT_CONVERSION_MODEL": "auto",
                "METRIKA_ATTRIBUTION_MODEL": "auto",
                "CLEAN_METRIKA_LOG_AFTER_DOWNLOAD": "false",
                "INCLUDE_VAT": "true",
                "INCLUDE_ASSOCIATED_REVENUE": "true",
            }
        )
        self.assertEqual(settings.metrika_counter_ids, (123, 456))
        self.assertEqual(settings.goal_ids, (99, 100))
        self.assertEqual(settings.lookback_days, 45)
        self.assertEqual(settings.direct_conversion_model, "AUTO")
        self.assertEqual(settings.metrika_attribution_model, "AUTOMATIC")
        self.assertFalse(settings.clean_metrika_log_after_download)
        self.assertTrue(settings.include_vat)
        self.assertTrue(settings.include_associated_revenue)

    def test_vat_is_included_by_default(self) -> None:
        settings = Settings.from_env(
            {
                "YANDEX_METRIKA_COUNTER_IDS": "123",
                "YANDEX_DIRECT_CLIENT_LOGIN": "client-login",
            }
        )
        self.assertTrue(settings.include_vat)
        self.assertFalse(settings.include_associated_revenue)
        self.assertEqual(settings.metrika_attribution_model, "LAST")

    def test_legacy_direct_cache_is_invalidated(self) -> None:
        settings = Settings(
            metrika_counter_ids=(1,),
            direct_client_login="client",
            goal_ids=(99,),
        )
        legacy_metadata = {
            "direct": {
                "client_login": "client",
                "attribution_model": "LC",
                "include_vat": True,
            }
        }
        self.assertFalse(direct_cache_matches(legacy_metadata, settings))
        legacy_metadata["direct"]["report_cache_version"] = (
            DIRECT_REPORT_CACHE_VERSION
        )
        self.assertTrue(direct_cache_matches(legacy_metadata, settings))

    def test_legacy_metrika_cache_is_invalidated(self) -> None:
        settings = Settings(
            metrika_counter_ids=(1,),
            direct_client_login="client",
            goal_ids=(99,),
        )
        legacy_metadata: dict[str, object] = {}
        self.assertFalse(metrika_cache_matches(legacy_metadata, settings))
        metrika_metadata = {
            "log_cache_version": METRIKA_LOG_CACHE_VERSION,
            "attribution_model": "AUTOMATIC",
            "lookback_days": 90,
        }
        legacy_metadata["metrika"] = metrika_metadata
        self.assertFalse(metrika_cache_matches(legacy_metadata, settings))
        metrika_metadata["attribution_model"] = "LAST"
        self.assertTrue(metrika_cache_matches(legacy_metadata, settings))
        revenue_settings = Settings(
            metrika_counter_ids=(1,),
            direct_client_login="client",
            goal_ids=(99,),
            include_associated_revenue=True,
        )
        self.assertFalse(metrika_cache_matches(legacy_metadata, revenue_settings))
        metrika_metadata["include_associated_revenue"] = True
        self.assertTrue(metrika_cache_matches(legacy_metadata, revenue_settings))

    def test_automatic_metrika_fields_are_loaded(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            path = root / "1" / "part_00000.tsv"
            path.parent.mkdir(parents=True)
            with path.open("w", encoding="utf-8", newline="") as fh:
                writer = csv.writer(fh, delimiter="\t")
                writer.writerow(
                    [
                        "ym:s:visitID",
                        "ym:s:clientID",
                        "ym:s:counterUserIDHash",
                        "ym:s:date",
                        "ym:s:dateTime",
                        "ym:s:goalsID",
                        "ym:s:automaticTrafficSource",
                        "ym:s:automaticDirectClickOrder",
                    ]
                )
                writer.writerow(
                    [
                        "1",
                        "10",
                        "100",
                        "2026-01-03",
                        "2026-01-03 10:00:00",
                        "[99]",
                        "ad",
                        "7",
                    ]
                )

            settings = Settings(
                metrika_counter_ids=(1,),
                direct_client_login="client",
                goal_ids=(99,),
                metrika_attribution_model="AUTOMATIC",
            )
            connection = connect(root / "test.sqlite3")
            create_schema(connection)
            counts = load_metrika_visits(connection, [path], settings)
            visit = connection.execute(
                "SELECT traffic_source, campaign_id FROM visits"
            ).fetchone()
            connection.close()

            self.assertEqual(counts, (1, 1))
            self.assertEqual(tuple(visit), ("yandex_direct", "7"))

    def test_direct_undefined_is_normalized_only_for_ad_traffic(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            path = root / "1" / "part_00000.tsv"
            path.parent.mkdir(parents=True)
            with path.open("w", encoding="utf-8", newline="") as fh:
                writer = csv.writer(fh, delimiter="\t")
                writer.writerow(
                    [
                        "ym:s:visitID",
                        "ym:s:clientID",
                        "ym:s:counterUserIDHash",
                        "ym:s:date",
                        "ym:s:dateTime",
                        "ym:s:goalsID",
                        "ym:s:automaticTrafficSource",
                        "ym:s:automaticAdvEngine",
                        "ym:s:automaticDirectClickOrder",
                    ]
                )
                writer.writerows(
                    [
                        [
                            "1", "10", "100", "2026-01-03",
                            "2026-01-03 10:00:00", "[]", "ad",
                            "ya_undefined", "",
                        ],
                        [
                            "2", "20", "200", "2026-01-04",
                            "2026-01-04 10:00:00", "[]", "organic",
                            "ya_undefined", "",
                        ],
                    ]
                )

            settings = Settings(
                metrika_counter_ids=(1,),
                direct_client_login="client",
                goal_ids=(99,),
                metrika_attribution_model="AUTOMATIC",
            )
            connection = connect(root / "test.sqlite3")
            create_schema(connection)
            counts = load_metrika_visits(connection, [path], settings)
            sources = connection.execute(
                "SELECT visit_id, traffic_source FROM visits ORDER BY visit_id"
            ).fetchall()
            connection.close()

            self.assertEqual(counts, (2, 0))
            self.assertEqual(
                [tuple(row) for row in sources],
                [("1", "yandex_direct"), ("2", "organic")],
            )

    def test_ua_style_assists_use_only_prior_direct_visits(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            settings = Settings(
                metrika_counter_ids=(1,),
                direct_client_login="client",
                goal_ids=(99,),
                lookback_days=90,
            )
            connection = connect(root / "test.sqlite3")
            create_schema(connection)

            visits = [
                (
                    1, "a", "user", "2026-07-01 10:00:00", "2026-07-01",
                    "yandex_direct", "Yandex Direct", "7", "Campaign A",
                    "70", "Group A", "alpha", "alpha", "KEYWORD",
                    "SEARCH", "yandex.ru", "yandex.ru", "", "", "", "", "",
                ),
                (
                    1, "b", "user", "2026-07-02 10:00:00", "2026-07-02",
                    "yandex_direct", "Yandex Direct", "9", "Campaign B",
                    "90", "Group B", "brand", "brand", "KEYWORD",
                    "SEARCH", "yandex.ru", "yandex.ru", "", "", "", "", "",
                ),
                (
                    1, "c1", "user", "2026-08-05 10:00:00", "2026-08-05",
                    "yandex_direct", "Yandex Direct", "9", "Campaign B",
                    "90", "Group B", "brand", "brand", "KEYWORD",
                    "SEARCH", "yandex.ru", "yandex.ru", "", "", "", "", "",
                ),
                (
                    1, "c2", "user", "2026-08-10 10:00:00", "2026-08-10",
                    "organic", "", "", "", "", "", "", "", "",
                    "", "", "", "", "", "", "", "",
                ),
                (
                    1, "d", "user", "2026-08-20 10:00:00", "2026-08-20",
                    "yandex_direct", "Yandex Direct", "11", "Campaign C",
                    "110", "Group C", "future", "future", "KEYWORD",
                    "SEARCH", "yandex.ru", "yandex.ru", "", "", "", "", "",
                ),
            ]
            connection.executemany(
                "INSERT INTO visits VALUES (" + ",".join("?" for _ in range(22)) + ")",
                visits,
            )
            connection.executemany(
                "INSERT INTO goal_visits VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)",
                [
                    (1, "c1", "user", "2026-08-05 10:00:00", "2026-08-05", 99, 1, 0, ""),
                    (1, "c2", "user", "2026-08-10 10:00:00", "2026-08-10", 99, 1, 0, ""),
                ],
            )
            connection.executemany(
                "INSERT INTO direct_stats VALUES ("
                + ",".join("?" for _ in range(17))
                + ")",
                [
                    (
                        "2026-08-01", "7", "Campaign A", "70", "Group A",
                        "700", "alpha", "alpha", "KEYWORD", "SEARCH",
                        "yandex.ru", "yandex.ru", 100, 10, 100, 0, 200,
                    ),
                    (
                        "2026-08-05", "9", "Campaign B", "90", "Group B",
                        "900", "brand", "brand", "KEYWORD", "SEARCH",
                        "yandex.ru", "yandex.ru", 100, 10, 100, 1, 400,
                    ),
                    (
                        "2026-08-20", "11", "Campaign C", "110", "Group C",
                        "1100", "future", "future", "KEYWORD", "SEARCH",
                        "yandex.ru", "yandex.ru", 100, 10, 100, 0, 0,
                    ),
                ],
            )
            connection.commit()

            self.assertEqual(
                build_conversion_touches(
                    connection, settings, date(2026, 8, 1), date(2026, 8, 31)
                ),
                2,
            )
            create_performance_tables(
                connection, settings, date(2026, 8, 1), date(2026, 8, 31)
            )

            source = connection.execute(
                """
                SELECT metrika_leads, assisted_metrika_conversions,
                       total_metrika_conversions
                FROM report_by_source
                WHERE source = 'yandex_direct'
                """
            ).fetchone()
            campaign_a = connection.execute(
                """
                SELECT metrika_leads, assisted_metrika_conversions
                FROM report_by_campaign WHERE campaign_id = '7'
                """
            ).fetchone()
            campaign_b = connection.execute(
                """
                SELECT metrika_leads, assisted_metrika_conversions,
                       direct_revenue, drr_direct_pct, romi_direct_pct
                FROM report_by_campaign WHERE campaign_id = '9'
                """
            ).fetchone()
            campaign_c = connection.execute(
                """
                SELECT metrika_leads, assisted_metrika_conversions
                FROM report_by_campaign WHERE campaign_id = '11'
                """
            ).fetchone()
            source_assist_rows = connection.execute(channel_assists_sql()).fetchall()
            source_statuses = connection.execute(
                f"""
                SELECT source_assist_status, COUNT(*)
                FROM ({conversion_touches_sql()})
                GROUP BY source_assist_status
                """
            ).fetchall()
            connection.close()

            self.assertEqual(tuple(source), (1, 1, 2))
            self.assertEqual(tuple(campaign_a), (0, 1))
            self.assertEqual(tuple(campaign_b), (1, 1, 400.0, 25.0, 300.0))
            self.assertEqual(tuple(campaign_c), (0, 0))
            self.assertEqual(sum(row["assisted_converted_visits"] for row in source_assist_rows), 2)
            self.assertEqual(
                {row[0]: row[1] for row in source_statuses},
                {"eligible_assist": 2},
            )

    def test_direct_undefined_assists_source_but_not_campaign(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            settings = Settings(
                metrika_counter_ids=(1,),
                direct_client_login="client",
                goal_ids=(99,),
                lookback_days=90,
            )
            connection = connect(root / "test.sqlite3")
            create_schema(connection)
            connection.executemany(
                "INSERT INTO visits VALUES (" + ",".join("?" for _ in range(22)) + ")",
                [
                    (
                        1, "d", "user", "2026-07-10 10:00:00", "2026-07-10",
                        "yandex_direct", "ya_undefined", "", "", "", "", "",
                        "", "", "", "", "", "", "", "", "", "",
                    ),
                    (
                        1, "c", "user", "2026-08-10 10:00:00", "2026-08-10",
                        "organic", "", "", "", "", "", "", "", "", "",
                        "", "", "", "", "", "", "",
                    ),
                ],
            )
            connection.execute(
                "INSERT INTO goal_visits VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)",
                (1, "c", "user", "2026-08-10 10:00:00", "2026-08-10", 99, 1, 0, ""),
            )
            connection.commit()

            self.assertEqual(
                build_conversion_touches(
                    connection, settings, date(2026, 8, 1), date(2026, 8, 31)
                ),
                1,
            )
            create_performance_tables(
                connection, settings, date(2026, 8, 1), date(2026, 8, 31)
            )
            source = connection.execute(
                """
                SELECT assisted_metrika_conversions
                FROM report_by_source
                WHERE source = 'yandex_direct'
                """
            ).fetchone()
            campaign_rows = connection.execute(
                "SELECT COUNT(*) FROM report_by_campaign"
            ).fetchone()[0]
            source_columns = {
                row[1]
                for row in connection.execute(
                    "PRAGMA table_info(report_by_source)"
                )
            }
            connection.close()

            self.assertEqual(source[0], 1)
            self.assertEqual(campaign_rows, 0)
            self.assertNotIn("assisted_metrika_revenue", source_columns)
            self.assertNotIn("goal_revenue", conversion_touches_sql())

    def test_one_direct_assist_and_one_last_touch_conversion(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            metrika = root / "metrika" / "1" / "part_00000.tsv"
            metrika_second = root / "metrika" / "2" / "part_00000.tsv"
            direct = root / "direct.tsv"
            settings = Settings(
                metrika_counter_ids=(1, 2),
                direct_client_login="client",
                goal_ids=(99, 100),
                lookback_days=30,
                include_associated_revenue=True,
            )

            metrika_fields = [
                "ym:s:visitID", "ym:s:clientID", "ym:s:counterUserIDHash",
                "ym:s:date", "ym:s:dateTime", "ym:s:goalsID",
                "ym:s:goalsPrice", "ym:s:goalsCurrency",
                "ym:s:lastTrafficSource", "ym:s:lastAdvEngine",
                "ym:s:lastDirectClickOrder", "ym:s:lastDirectBannerGroup",
                "ym:s:lastDirectClickOrderName", "ym:s:lastClickBannerGroupName",
                "ym:s:lastDirectPhraseOrCond", "ym:s:lastDirectPlatformType",
                "ym:s:lastDirectPlatform", "ym:s:lastDirectConditionType",
                "ym:s:lastUTMSource", "ym:s:lastUTMMedium",
                "ym:s:lastUTMCampaign", "ym:s:lastUTMContent", "ym:s:lastUTMTerm",
            ]
            rows = [
                ["1", "10", "100", "2026-01-01", "2026-01-01 10:00:00", "[]",
                 "[]", "[]",
                 "ad", "Yandex Direct", "7", "8", "Campaign", "Group", "red shoes",
                 "Search", "yandex.ru", "KEYWORD", "", "", "", "", ""],
                ["3", "10", "100", "2026-01-02", "2026-01-02 10:00:00", "[]",
                 "[]", "[]",
                 "ad", "Yandex Direct", "7", "8", "Campaign", "Group", "blue shoes",
                 "Search", "yandex.ru", "KEYWORD", "", "", "", "", ""],
                ["2", "10", "100", "2026-01-03", "2026-01-03 10:00:00", "[99,100]",
                 "[100000,250000]", "['RUB','RUB']",
                 "ad", "Yandex Direct", "9", "10", "Last", "Last Group", "brand",
                 "Search", "yandex.ru", "KEYWORD", "", "", "", "", ""],
                ["4", "10", "100", "2026-01-04", "2026-01-04 10:00:00", "[99]",
                 "[500000]", "['RUB']",
                 "organic", "", "", "", "", "", "", "", "", "", "", "", "", "", ""],
            ]
            metrika.parent.mkdir(parents=True)
            metrika_second.parent.mkdir(parents=True)
            with metrika.open("w", encoding="utf-8", newline="") as fh:
                writer = csv.writer(fh, delimiter="\t")
                writer.writerow(metrika_fields)
                writer.writerows(rows)
            with metrika_second.open("w", encoding="utf-8", newline="") as fh:
                writer = csv.writer(fh, delimiter="\t")
                writer.writerow(metrika_fields)
                writer.writerows(rows)

            direct_fields = [
                "Date", "CampaignId", "CampaignName", "AdGroupId", "AdGroupName",
                "CriterionId", "Criterion", "CriterionType", "AdNetworkType",
                "Placement", "Impressions", "Clicks", "Cost",
                "Conversions_99_LC", "Conversions_100_LC",
                "Revenue_99_LC", "Revenue_100_LC",
            ]
            with direct.open("w", encoding="utf-8", newline="") as fh:
                writer = csv.writer(fh, delimiter="\t")
                writer.writerow(direct_fields)
                writer.writerow([
                    "2026-01-01", "7", "Campaign", "8", "Group", "11", "red shoes",
                    "KEYWORD", "SEARCH", "yandex.ru", "100", "10", "123.45", "0", "0", "0", "0",
                ])
                writer.writerow([
                    "2026-01-03", "9", "Last", "10", "Last Group", "12", "brand",
                    "KEYWORD", "SEARCH", "yandex.ru", "50", "5", "50", "1", "1", "100", "200",
                ])

            qa = build(
                [direct], [metrika, metrika_second],
                root / "db.sqlite3", root / "report.xlsx", settings,
                date(2026, 1, 1), date(2026, 1, 31),
            )
            self.assertEqual(qa["metrika_counter_ids"], [1, 2])
            self.assertEqual(qa["loaded"]["conversion_touches"], 6)
            self.assertEqual(qa["combined_totals"]["direct_impressions"], 150.0)
            self.assertEqual(qa["combined_totals"]["direct_revenue"], 300.0)
            self.assertAlmostEqual(
                qa["combined_totals"]["direct_drr_pct"], 57.816667
            )
            self.assertAlmostEqual(
                qa["combined_totals"]["direct_romi_pct"], 72.960507
            )
            self.assertEqual(qa["combined_totals"]["unique_goal_visits"], 4)
            self.assertEqual(qa["combined_totals"]["goal_reaches"], 6)
            self.assertEqual(
                qa["combined_totals"]["assisted_metrika_conversions_unique"],
                2,
            )
            self.assertEqual(
                qa["combined_totals"]["metrika_conversions_with_assists"],
                6,
            )
            self.assertEqual(
                qa["combined_totals"]["assisted_metrika_revenue_unique"],
                1000.0,
            )

            workbook = load_workbook(root / "report.xlsx", read_only=True, data_only=True)
            worksheet = workbook["Полный срез"]
            headers = [cell.value for cell in next(worksheet.iter_rows())]
            output = [
                dict(zip(headers, row))
                for row in worksheet.iter_rows(min_row=2, values_only=True)
            ]
            assist = next(
                row
                for row in output
                if row["ID кампании"] == "7"
                and row["Ключ / условие показа"] == "blue shoes"
            )
            conversion = next(row for row in output if row["ID кампании"] == "9")
            self.assertEqual(assist["Цели Метрики"], "99,100")
            self.assertEqual(assist["Счётчики Метрики"], "1,2")
            self.assertEqual(assist["Ассоциированные конверсии Метрики"], 2)
            self.assertEqual(assist["Ассоциированный доход Метрики"], 1000.0)
            self.assertEqual(conversion["Лиды Метрики"], 2)
            self.assertEqual(conversion["Ассоциированные конверсии Метрики"], 2)
            self.assertEqual(conversion["Лиды Директа"], 2.0)
            self.assertEqual(conversion["Модель атрибуции Директа"], "LC — последний переход")
            self.assertEqual(
                conversion["Модель атрибуции Метрики"],
                "LAST — последний переход",
            )
            self.assertEqual(conversion["CR Директа, %"], 40.0)
            self.assertEqual(conversion["CPL Директа"], 25.0)
            self.assertEqual(conversion["Доход Директа"], 300.0)
            self.assertEqual(conversion["ДРР Директа, %"], 16.6667)
            self.assertEqual(conversion["ROMI Директа, %"], 500.0)
            self.assertIn("Расход с НДС", headers)
            self.assertTrue(
                qa["conversion_cohort"]["direct_cost_includes_vat"]
            )
            self.assertNotIn("CR ассоциированных конверсий, %", headers)
            self.assertNotIn("CPL ассоциированных конверсий", headers)
            removed_columns = {
                "Ассоциированные конверсии Метрики (линейно)",
                "Ассоциированные достижения (линейно)",
                "Ассоциированный доход Метрики (линейно)",
            }
            self.assertTrue(removed_columns.isdisjoint(headers))
            self.assertIn("Обычные достижения целей", headers)
            self.assertIn("Ассоциированные достижения целей", headers)
            self.assertIn("Ассоциированный доход Метрики", headers)
            self.assertEqual(conversion["CR лидов Метрики, %"], 40.0)
            self.assertEqual(conversion["CPL лидов Метрики"], 25.0)
            self.assertEqual(conversion["Статус ассоциированности"], "Есть лиды Метрики")
            self.assertIn("Кампании", workbook.sheetnames)
            self.assertIn("Площадки", workbook.sheetnames)
            self.assertIn("Основные пути", workbook.sheetnames)
            self.assertIn("Роль каналов", workbook.sheetnames)
            self.assertEqual(qa["outputs"]["channel_roles_rows"], 2)
            self.assertEqual(workbook.sheetnames[1], "Справка")
            help_sheet = workbook["Справка"]
            help_headers = [cell.value for cell in next(help_sheet.iter_rows())]
            self.assertEqual(
                help_headers,
                [
                    "Раздел",
                    "Столбец",
                    "Что означает",
                    "Формула / важная особенность",
                ],
            )
            help_rows = list(help_sheet.iter_rows(min_row=2, values_only=True))
            assisted_help = next(
                row
                for row in help_rows
                if row[1] == "Ассоциированные конверсии Метрики"
            )
            self.assertIn("неаддитивен", assisted_help[3])
            self.assertTrue(
                any(row[1] == "Расход с НДС" for row in help_rows)
            )
            self.assertTrue(removed_columns.isdisjoint(row[1] for row in help_rows))
            self.assertTrue(any(row[1] == "Доход Директа" for row in help_rows))
            self.assertTrue(any(row[1] == "ДРР Директа, %" for row in help_rows))
            self.assertTrue(any(row[1] == "ROMI Директа, %" for row in help_rows))
            source_sheet = workbook["Источники"]
            source_headers = [cell.value for cell in next(source_sheet.iter_rows())]
            by_source = [
                dict(zip(source_headers, row))
                for row in source_sheet.iter_rows(min_row=2, values_only=True)
            ]
            yandex = next(row for row in by_source if row["Источник"] == "yandex_direct")
            self.assertEqual(yandex["Коэффициент ассоциированности"], 1.0)
            self.assertEqual(yandex["Всего конверсий Метрики"], 4)
            self.assertEqual(yandex["Ассоциированный доход Метрики"], 1000.0)
            self.assertEqual(yandex["Доход Директа"], 300.0)
            self.assertEqual(yandex["ДРР Директа, %"], 57.8167)
            self.assertEqual(yandex["ROMI Директа, %"], 72.9605)
            self.assertEqual(yandex["CR всех конверсий Метрики, %"], 26.6667)
            self.assertEqual(yandex["CPL всех конверсий Метрики"], 43.36)
            cr_column = source_headers.index("CR лидов Метрики, %") + 1
            self.assertEqual(
                source_sheet.cell(2, cr_column).number_format,
                '0.0"%"',
            )
            drr_column = source_headers.index("ДРР Директа, %") + 1
            self.assertEqual(
                source_sheet.cell(2, drr_column).number_format,
                '0.0"%"',
            )
            campaign_sheet = workbook["Кампании"]
            campaign_headers = [
                cell.value for cell in next(campaign_sheet.iter_rows())
            ]
            campaigns = [
                dict(zip(campaign_headers, row))
                for row in campaign_sheet.iter_rows(min_row=2, values_only=True)
            ]
            assisted_campaign = next(
                row for row in campaigns if row["ID кампании"] == "7"
            )
            self.assertEqual(
                assisted_campaign["Ассоциированные конверсии Метрики"],
                2,
            )
            path_sheet = workbook["Основные пути"]
            path_headers = [
                cell.value for cell in next(path_sheet.iter_rows())
            ]
            paths = [
                dict(zip(path_headers, row))
                for row in path_sheet.iter_rows(min_row=2, values_only=True)
            ]
            self.assertEqual(
                sum(row["Конверсионные визиты"] for row in paths),
                4,
            )
            assisted_path = next(
                row for row in paths
                if row["Путь конверсии"].endswith("Поисковые системы")
            )
            self.assertEqual(
                assisted_path["Путь конверсии"],
                "Яндекс Директ → Поисковые системы",
            )
            self.assertEqual(assisted_path["Конверсионные визиты"], 2)
            role_sheet = workbook["Роль каналов"]
            role_headers = [cell.value for cell in next(role_sheet.iter_rows())]
            self.assertEqual(len(role_headers), 8)
            roles = [
                dict(zip(role_headers, row))
                for row in role_sheet.iter_rows(min_row=2, values_only=True)
            ]
            direct_role = next(row for row in roles if row["Источник"] == "Яндекс Директ")
            self.assertEqual(direct_role["Средний вес в цепочке, %"], 85.7143)
            self.assertEqual(direct_role["Доля уникальных цепочек, %"], 100.0)
            self.assertEqual(direct_role["Доля промежуточных касаний, %"], 100.0)
            self.assertEqual(direct_role["Доля последних касаний, %"], 50.0)
            self.assertEqual(role_sheet.cell(2, 2).number_format, '0.0"%"')
            self.assertEqual(role_sheet.cell(2, 2).data_type, "n")
            for role_header in role_headers:
                self.assertTrue(any(row[0] == "Роль каналов" and row[1] == role_header for row in help_rows))
            for sheet in workbook.worksheets:
                values = {
                    value
                    for row in sheet.iter_rows(values_only=True)
                    for value in row
                    if isinstance(value, str)
                }
                self.assertTrue(
                    removed_columns.isdisjoint(values),
                    f"Удалённый столбец найден на листе {sheet.title}",
                )
            workbook.close()
            json.loads(json.dumps(qa))



if __name__ == "__main__":
    unittest.main()
