from __future__ import annotations

import re
import sqlite3
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
BODY_FONT = Font(name="Arial", size=10, color="222222")
TITLE_FONT = Font(name="Arial", size=15, bold=True, color="1F1F1F")
SECTION_FONT = Font(name="Arial", size=11, bold=True, color="1F4E78")
HEADER_BORDER = Border(
    bottom=Side(style="thin", color="FFFFFF"),
    right=Side(style="thin", color="FFFFFF"),
)


COLUMN_LABELS = {
    "date_from": "Дата начала",
    "date_to": "Дата окончания",
    "counter_ids": "Счётчики Метрики",
    "goal_ids": "Цели Метрики",
    "lead_basis": "База лидов Метрики",
    "direct_attribution_model": "Модель атрибуции Директа",
    "metrika_attribution_model": "Модель атрибуции Метрики",
    "source": "Источник",
    "impressions": "Показы",
    "clicks": "Клики",
    "cost": "Расход",
    "direct_revenue": "Доход Директа",
    "drr_direct_pct": "ДРР Директа, %",
    "romi_direct_pct": "ROMI Директа, %",
    "direct_api_converted_visits": "Конверсии Директа (контроль)",
    "direct_leads": "Лиды Директа",
    "cr_direct_pct": "CR Директа, %",
    "cpl_direct": "CPL Директа",
    "ordinary_converted_visits": "Обычные конверсионные визиты",
    "ordinary_goal_reaches": "Обычные достижения целей",
    "assisted_converted_visits": "Ассоциированные конверсионные визиты",
    "assisted_goal_reaches": "Ассоциированные достижения целей",
    "assisted_conversions_linear": "Ассоциированные конверсии (линейно)",
    "metrika_leads": "Лиды Метрики",
    "assisted_metrika_conversions": "Ассоциированные конверсии Метрики",
    "assisted_metrika_revenue": "Ассоциированный доход Метрики",
    "total_metrika_conversions": "Всего конверсий Метрики",
    "cr_metrika_pct": "CR лидов Метрики, %",
    "cpl_metrika": "CPL лидов Метрики",
    "cr_total_metrika_pct": "CR всех конверсий Метрики, %",
    "cpl_total_metrika": "CPL всех конверсий Метрики",
    "association_coefficient": "Коэффициент ассоциированности",
    "association_status": "Статус ассоциированности",
    "campaign_id": "ID кампании",
    "campaign_name": "Кампания",
    "ad_group_id": "ID группы",
    "ad_group_name": "Группа объявлений",
    "criterion_id": "ID условия показа",
    "criterion": "Ключ / условие показа",
    "criterion_type": "Тип условия показа",
    "placement_type": "Тип площадки",
    "placement": "Площадка",
    "counter_id": "Счётчик Метрики",
    "adv_engine": "Рекламная система",
    "utm_source": "UTM-источник",
    "utm_medium": "UTM-канал",
    "utm_campaign": "UTM-кампания",
    "utm_content": "UTM-содержание",
    "utm_term": "UTM-термин",
    "conversion_visit_id": "ID конверсионного визита",
    "goal_reaches": "Достижения целей",
    "goal_revenue": "Доход выбранных целей",
    "conversion_date": "Дата конверсии",
    "closing_source": "Источник конверсионного визита",
    "source_assist_status": "Ассист по источнику",
    "touch_visit_id": "ID визита-касания",
    "touch_datetime": "Дата и время касания",
    "touch_date": "Дата касания",
    "linear_conversion_credit": "Линейный кредит конверсии",
    "linear_goal_reach_credit": "Линейный кредит достижений",
    "linear_revenue_credit": "Линейный кредит дохода",
    "conversion_path": "Путь конверсии",
    "path_length": "Шагов в пути",
    "closing_channel": "Закрывающий канал",
    "contains_yandex_direct": "Яндекс Директ в пути",
    "converted_visits": "Конверсионные визиты",
    "conversion_share_pct": "Доля конверсий, %",
    "chain_weight_pct": "Средний вес в цепочке, %",
    "unique_chain_share_pct": "Доля уникальных цепочек, %",
    "first_touch_share_pct": "Доля первых касаний, %",
    "first_touch_2plus_share_pct": "Доля первых касаний (2+ касания), %",
    "middle_touch_share_pct": "Доля промежуточных касаний, %",
    "last_touch_share_pct": "Доля последних касаний, %",
    "last_touch_2plus_share_pct": "Доля последних касаний (2+ касания), %",
}


SHEET_SPECS = (
    ("Источники", "SELECT * FROM report_by_source"),
    ("Кампании", "SELECT * FROM report_by_campaign"),
    ("Группы", "SELECT * FROM report_by_ad_group"),
    ("Ключи", "SELECT * FROM report_by_criterion"),
    ("Типы площадок", "SELECT * FROM report_by_placement_type"),
    ("Площадки", "SELECT * FROM report_by_placement"),
    ("Полный срез", "SELECT * FROM report_by_full_detail"),
    ("Основные пути", "SELECT * FROM report_conversion_paths"),
    ("Роль каналов", "SELECT * FROM report_channel_roles ORDER BY unique_chain_share_pct DESC, chain_weight_pct DESC, source"),
)


def _safe_value(value: object, column_name: str) -> object:
    if value is None:
        return None
    if isinstance(value, str):
        cleaned = ILLEGAL_CHARACTERS_RE.sub("", value)
        if len(cleaned) > 32767:
            cleaned = cleaned[:32767]
        if column_name in {"date_from", "date_to", "conversion_date", "touch_date"}:
            try:
                return date.fromisoformat(cleaned)
            except ValueError:
                return cleaned
        if column_name.endswith("datetime"):
            try:
                return datetime.fromisoformat(cleaned)
            except ValueError:
                return cleaned
        if column_name == "lead_basis" and cleaned == "converted_visits":
            return "Уникальные визиты с целью"
        if column_name == "direct_attribution_model":
            model_labels = {
                "LC": "LC — последний переход",
                "LSCCD": "LSCCD — последний значимый переход, кросс-девайс",
                "FCCD": "FCCD — первый переход, кросс-девайс",
                "AUTO": "AUTO — автоматическая атрибуция",
            }
            return model_labels.get(cleaned, cleaned)
        if column_name == "metrika_attribution_model":
            model_labels = {
                "AUTOMATIC": "AUTO — автоматическая атрибуция",
                "LAST": "LAST — последний переход",
                "LASTSIGN": "LASTSIGN — последний значимый переход",
                "FIRST": "FIRST — первый переход",
                "LAST_YANDEX_DIRECT_CLICK": (
                    "Последний переход из Яндекс Директа"
                ),
                "CROSS_DEVICE_LAST_SIGNIFICANT": (
                    "Последний значимый переход, кросс-девайс"
                ),
                "CROSS_DEVICE_FIRST": "Первый переход, кросс-девайс",
                "CROSS_DEVICE_LAST_YANDEX_DIRECT_CLICK": (
                    "Последний переход из Яндекс Директа, кросс-девайс"
                ),
                "CROSS_DEVICE_LAST": "Последний переход, кросс-девайс",
            }
            return model_labels.get(cleaned, cleaned)
        if column_name == "association_status":
            status_labels = {
                "ok": "Есть лиды Метрики",
                "only_assisted": "Только ассоциированные",
                "no_leads": "Конверсий нет",
            }
            return status_labels.get(cleaned, cleaned)
        if column_name == "source_assist_status":
            status_labels = {
                "eligible_assist": "Да",
                "same_closing_source": "Нет — источник закрыл конверсию",
            }
            return status_labels.get(cleaned, cleaned)
        return cleaned
    return value


def _number_format(column_name: str) -> str | None:
    if column_name in {"date_from", "date_to", "conversion_date", "touch_date"}:
        return "yyyy-mm-dd"
    if column_name.endswith("datetime"):
        return "yyyy-mm-dd hh:mm:ss"
    if (
        column_name == "cost"
        or column_name.startswith("cpl_")
        or "revenue" in column_name
    ):
        return "#,##0.0"
    if column_name.endswith("_pct"):
        return '0.0"%"'
    if column_name in {"association_coefficient"}:
        return "0.0"
    if column_name.endswith("_id"):
        return "0"
    if any(
        marker in column_name
        for marker in (
            "impressions",
            "clicks",
            "leads",
            "visits",
            "reaches",
            "conversions",
            "credit",
        )
    ):
        return "#,##0.0"
    return None


def _column_width(column_name: str) -> float:
    if column_name == "conversion_path":
        return 60
    if column_name in {"campaign_name", "ad_group_name", "criterion", "placement"}:
        return 30
    if column_name.startswith("utm_") or column_name in {"adv_engine", "source"}:
        return 22
    if column_name in {
        "association_status",
        "lead_basis",
        "criterion_type",
        "direct_attribution_model",
        "metrika_attribution_model",
    }:
        return 20
    if column_name.endswith("datetime"):
        return 20
    if "date" in column_name:
        return 13
    if column_name.endswith("_id") or column_name.endswith("_ids"):
        return 18
    label = COLUMN_LABELS.get(column_name, column_name)
    if len(label) >= 24:
        return 24
    return max(12, min(22, len(label) + 2))


def _header_cell(worksheet, value: str) -> WriteOnlyCell:
    cell = WriteOnlyCell(worksheet, value=value)
    cell.font = HEADER_FONT
    cell.fill = HEADER_FILL
    cell.alignment = Alignment(
        horizontal="center", vertical="center", wrap_text=True
    )
    cell.border = HEADER_BORDER
    return cell


def _body_cell(worksheet, value: object, column_name: str) -> WriteOnlyCell:
    cell = WriteOnlyCell(worksheet, value=_safe_value(value, column_name))
    cell.font = BODY_FONT
    cell.alignment = Alignment(vertical="center")
    number_format = _number_format(column_name)
    if number_format:
        cell.number_format = number_format
    elif isinstance(value, (int, float)) and not isinstance(value, bool):
        cell.number_format = "#,##0.0"
    return cell


def _add_query_sheet(
    workbook: Workbook,
    connection: sqlite3.Connection,
    sheet_name: str,
    sql: str,
    cost_includes_vat: bool,
) -> int:
    worksheet = workbook.create_sheet(sheet_name)
    worksheet.sheet_view.showGridLines = False
    worksheet.sheet_properties.tabColor = "4472C4"
    worksheet.freeze_panes = "A2"
    worksheet.row_dimensions[1].height = 60 if sheet_name == "Роль каналов" else 45
    if sheet_name == "Роль каналов":
        worksheet.freeze_panes = "B2"

    cursor = connection.execute(sql)
    headers = [item[0] for item in cursor.description]
    for index, header in enumerate(headers, start=1):
        worksheet.column_dimensions[get_column_letter(index)].width = _column_width(
            header
        )
    if sheet_name == "Роль каналов":
        worksheet.column_dimensions["A"].width = 34
    worksheet.append(
        [
            _header_cell(
                worksheet,
                (
                    "Расход с НДС"
                    if header == "cost" and cost_includes_vat
                    else "Расход без НДС"
                    if header == "cost"
                    else COLUMN_LABELS.get(header, header)
                ),
            )
            for header in headers
        ]
    )

    row_count = 0
    for row in cursor:
        worksheet.append(
            [
                _body_cell(worksheet, value, headers[index])
                for index, value in enumerate(row)
            ]
        )
        row_count += 1

    last_column = get_column_letter(max(1, len(headers)))
    worksheet.auto_filter.ref = f"A1:{last_column}{row_count + 1}"
    return row_count


def _column_help_rows(
    cost_includes_vat: bool,
    include_associated_revenue: bool,
) -> list[tuple[str, str, str, str]]:
    cost_label = "Расход с НДС" if cost_includes_vat else "Расход без НДС"
    rows = [
        ("Параметры", "Дата начала", "Первый день отчётного периода.", "Дата включается в расчёт."),
        ("Параметры", "Дата окончания", "Последний день отчётного периода.", "Дата включается в расчёт."),
        ("Параметры", "Счётчики Метрики", "ID объединённых счётчиков Метрики.", "Визиты разных счётчиков изолируются до агрегации."),
        ("Параметры", "Цели Метрики", "ID выбранных конечных целей.", "Все цели объединяются до расчёта KPI."),
        ("Параметры", "База лидов Метрики", "Единица расчёта обычных лидов Метрики.", "Один визит с одной или несколькими выбранными целями считается одним лидом."),
        ("Параметры", "Модель атрибуции Директа", "Модель из DIRECT_CONVERSION_MODEL.", "Применяется только к лидам, которые возвращает Reports API Директа."),
        ("Параметры", "Модель атрибуции Метрики", "Модель из METRIKA_ATTRIBUTION_MODEL.", "Для логики в стиле Universal Analytics используется LAST: он сохраняет фактический источник каждого визита."),
        ("Основные показатели", "Показы", "Количество показов из Reports API Директа.", "Аддитивный показатель."),
        ("Основные показатели", "Клики", "Количество кликов из Reports API Директа.", "Знаменатель для CR."),
        ("Основные показатели", cost_label, "Расход из Reports API Директа.", "Параметр IncludeVAT передаётся в API. Все CPL используют этот расход."),
        ("Основные показатели", "Доход Директа", "Сумма ценностей конверсий выбранных целей из Reports API Директа.", "Суммируются поля Revenue_<goal>_<model> по модели DIRECT_CONVERSION_MODEL. Это ценность целей, а не обязательно ecommerce-выручка."),
        ("Основные показатели", "ДРР Директа, %", "Доля рекламных расходов в доходе Директа.", "Расход / доход Директа × 100%. Пусто, если доход равен нулю."),
        ("Основные показатели", "ROMI Директа, %", "Окупаемость рекламных расходов по доходу Директа.", "(Доход Директа − расход) / расход × 100%. Пусто, если расход равен нулю; себестоимость и другие затраты не учитываются."),
        ("Основные показатели", "Лиды Директа", "Сумма конверсий выбранных целей по модели атрибуции Директа.", "Один визит может попасть в несколько полей целей, поэтому показатель может превышать уникальные лиды Метрики."),
        ("Основные показатели", "CR Директа, %", "Доля лидов Директа от кликов.", "Лиды Директа / клики × 100%."),
        ("Основные показатели", "CPL Директа", "Стоимость одного лида Директа.", "Расход / лиды Директа."),
        ("Основные показатели", "Лиды Метрики", "Уникальные визиты с хотя бы одной выбранной целью.", "Распределяются по источнику конверсионного визита согласно модели атрибуции Метрики."),
        ("Основные показатели", "CR лидов Метрики, %", "Доля обычных лидов Метрики от кликов.", "Лиды Метрики / клики × 100%."),
        ("Основные показатели", "CPL лидов Метрики", "Стоимость обычного лида Метрики.", "Расход / лиды Метрики."),
        ("Основные показатели", "Ассоциированные конверсии Метрики", "Уникальные целевые визиты не из Директа, перед которыми у того же посетителя был визит из Директа в пределах LOOKBACK_DAYS.", "Окно считается назад от каждой конверсии; будущие визиты исключены. «Директ: Не определено» учитывается на уровне источника. В строке конверсия учитывается один раз, но между строками показатель неаддитивен."),
        ("Основные показатели", "Всего конверсий Метрики", "Аналитическая сумма обычных и ассоциированных конверсий.", "Лиды Метрики + ассоциированные конверсии Метрики. Внутри одной строки одна конверсия не может быть и обычной, и ассоциированной."),
        ("Основные показатели", "CR всех конверсий Метрики, %", "Доля всех конверсий Метрики от кликов.", "Всего конверсий Метрики / клики × 100%."),
        ("Основные показатели", "CPL всех конверсий Метрики", "Стоимость конверсии с учётом ассистов.", "Расход / всего конверсий Метрики."),
        ("Основные показатели", "Коэффициент ассоциированности", "Отношение ассистирующих конверсий к обычным лидам.", "Ассоциированные конверсии Метрики / лиды Метрики."),
        ("Основные показатели", "Статус ассоциированности", "Подсказка для строк с обычными и/или ассоциированными лидами.", "«Только ассоциированные» означает, что знаменатель коэффициента равен нулю."),
        ("Диагностика", "Обычные достижения целей", "Общее число срабатываний выбранных целей в закрывающих визитах.", "Если цель сработала несколько раз или визит достиг нескольких целей, учитывается несколько достижений."),
        ("Диагностика", "Ассоциированные достижения целей", "Число достижений целей, полностью приписанное каждой ассистирующей строке.", "Неаддитивно между строками: одни достижения могут повторяться у нескольких касаний."),
        ("Срезы", "Источник", "Источник трафика визита.", "Для расходов Директа используется yandex_direct."),
        ("Срезы", "ID кампании / Кампания", "Идентификатор и название кампании.", "Пусто у источников без данных Директа."),
        ("Срезы", "ID группы / Группа объявлений", "Идентификатор и название группы объявлений.", "Пусто, если измерение не определено."),
        ("Срезы", "ID условия показа", "ID критерия показа из Директа.", "Для визитов Метрики надёжный ID доступен не всегда."),
        ("Срезы", "Ключ / условие показа", "Ключевая фраза или другое условие показа.", "Может обозначать автотаргетинг, ретаргетинг или условие фида."),
        ("Срезы", "Тип условия показа", "Тип критерия Директа.", "Помогает отличать ключевые слова от других условий."),
        ("Срезы", "Тип площадки", "Поиск или рекламная сеть.", "Нормализуется в SEARCH или AD_NETWORK, когда тип распознан."),
        ("Срезы", "Площадка", "Домен или название площадки показа.", "На поиске обычно yandex.ru; в сетях — площадка РСЯ."),
        ("Технические поля", "Счётчик Метрики", "Счётчик конкретного визита или конверсии.", "Сохраняется на технических листах для проверки источника данных."),
        ("Технические поля", "Рекламная система", "Рекламный движок, определённый Метрикой.", "Может отличаться от общего источника трафика."),
        ("Технические поля", "UTM-источник / UTM-канал / UTM-кампания", "Поля utm_source, utm_medium и utm_campaign визита.", "Используются на листах ассистов и цепочек."),
        ("Технические поля", "UTM-содержание / UTM-термин", "Поля utm_content и utm_term визита.", "Используются для дополнительной детализации касаний."),
        ("Технические поля", "Ассоциированные конверсионные визиты", "Количество строк-конверсий, которым набор измерений ассистировал.", "На техническом листе показатель не схлопывается до логики KPI каждого аналитического среза."),
        ("Технические поля", "ID конверсионного визита", "Идентификатор визита, в котором достигнута выбранная цель.", "Уникален только вместе с ID счётчика."),
        ("Технические поля", "Достижения целей", "Количество выбранных целей и их повторных достижений в конверсионном визите.", "Не равно уникальным лидам."),
        ("Технические поля", "Дата конверсии", "Дата закрывающего визита с целью.", "Относится к отчётному периоду."),
        ("Технические поля", "Источник конверсионного визита", "Источник, которому модель Метрики присвоила обычную конверсию.", "Для ассоциированной конверсии источник закрывающего визита не должен относиться к Директу."),
        ("Технические поля", "Ассист по источнику", "Показывает, прошёл ли целевой визит условие ассоциированности на уровне источника.", "В текущей выгрузке сохраняются только допустимые связи с Директом."),
        ("Технические поля", "ID визита-касания", "Идентификатор предыдущего визита из Директа.", "Визит всегда раньше конверсионного визита и не дальше LOOKBACK_DAYS до него."),
        ("Технические поля", "Дата и время касания / Дата касания", "Момент предыдущего визита из Директа.", "Окно LOOKBACK_DAYS отсчитывается назад от даты и времени каждой конверсии."),
        ("Технические поля", "Линейный кредит конверсии", "Доля одной конверсии между найденными визитами Директа.", "Техническое поле цепочки."),
        ("Технические поля", "Линейный кредит достижений", "Доля достижений целей между найденными визитами Директа.", "Техническое поле цепочки."),
        ("Пути конверсий", "Путь конверсии", "Последовательность каналов за LOOKBACK_DAYS до целевого визита, включая закрывающий канал.", "Одинаковые соседние каналы схлопываются; порядок слева направо — от раннего касания к конверсии."),
        ("Пути конверсий", "Шагов в пути", "Число канальных шагов после схлопывания повторов.", "Например, «Директ → Директ → Поиск» превращается в «Директ → Поиск» и имеет 2 шага."),
        ("Пути конверсий", "Закрывающий канал", "Последний канал пути, в визите которого достигнута цель.", "Этот канал получает обычный лид Метрики."),
        ("Пути конверсий", "Яндекс Директ в пути", "Признак наличия Директа на любом шаге пути.", "Значения: «Да» или «Нет»."),
        ("Пути конверсий", "Конверсионные визиты", "Число уникальных целевых визитов с данным путём.", "Сумма по листу равна общему числу лидов Метрики."),
        ("Пути конверсий", "Доля конверсий, %", "Доля целевых визитов, пришедшихся на данный путь.", "Конверсионные визиты пути / все конверсионные визиты × 100%."),
        ("Роль каналов", "Источник", "Канал визита, как на листе «Основные пути». Поисковые системы объединены, без разделения Google и Яндекса.", "Одна цепочка на каждый целевой визит, история внутри счётчика за LOOKBACK_DAYS до него плюс закрывающий визит. Повторные визиты не схлопываются."),
        ("Роль каналов", "Средний вес в цепочке, %", "Доля касаний канала среди всех касаний в конверсионных цепочках.", "Касания канала / все касания × 100%. Это общая доля, не среднее арифметическое долей отдельных цепочек. Сумма по каналам = 100%."),
        ("Роль каналов", "Доля уникальных цепочек, %", "В какой доле цепочек встретился канал хотя бы один раз.", "Цепочки с каналом / все цепочки × 100%. Повторы канала внутри цепочки не умножают её. Одинаковый путь к разным целевым визитам — разные цепочки."),
        ("Роль каналов", "Доля первых касаний, %", "Доля всех цепочек, в которых канал стоит первым.", "Цепочки с каналом в начале / все цепочки × 100%. Одновизитные цепочки включены."),
        ("Роль каналов", "Доля первых касаний (2+ касания), %", "Доля первых касаний среди цепочек как минимум из двух визитов.", "Цепочки из 2+ визитов с каналом в начале / все цепочки из 2+ визитов × 100%."),
        ("Роль каналов", "Доля промежуточных касаний, %", "Доля цепочек из 3+ визитов, где канал есть между первым и последним визитом.", "Цепочки с каналом в середине / все цепочки из 3+ визитов × 100%. Канал считается один раз в цепочке. Он также может быть первым или последним."),
        ("Роль каналов", "Доля последних касаний, %", "Доля всех цепочек, в которых канал закрыл целевой визит.", "Цепочки с каналом в конце / все цепочки × 100%. Одновизитные цепочки включены."),
        ("Роль каналов", "Доля последних касаний (2+ касания), %", "Доля последних касаний среди цепочек как минимум из двух визитов.", "Цепочки из 2+ визитов с каналом в конце / все цепочки из 2+ визитов × 100%."),
        ("Роль каналов", "Правила сравнения", "Проценты относятся к конверсионным цепочкам, а не ко всем посетителям сайта и не только к ассистам Директа.", "Без подходящих цепочек доля пуста. Доли уникальных цепочек и промежуточных касаний по каналам могут суммарно превышать 100%."),
    ]
    if include_associated_revenue:
        first_slice_row = next(
            index for index, row in enumerate(rows) if row[0] == "Срезы"
        )
        rows.insert(
            first_slice_row,
            (
                "Основные показатели",
                "Ассоциированный доход Метрики",
                "Сумма дохода выбранных целей в ассоциированных целевых визитах.",
                "Доход целевого визита полностью присваивается каждой допустимой строке среза, поэтому между строками показатель неаддитивен.",
            ),
        )
        rows.append(
            (
                "Технические поля",
                "Доход выбранных целей",
                "Сумма дохода выбранных целей в конверсионном визите.",
                "Берётся из ym:s:goalsPrice и делится на коэффициент Logs API 1000.",
            )
        )
        rows.append(
            (
                "Технические поля",
                "Линейный кредит дохода",
                "Доля дохода конверсионного визита на конкретное касание.",
                "Доход визита / число найденных касаний Директа.",
            )
        )
    return rows


def _add_help_sheet(
    workbook: Workbook,
    cost_includes_vat: bool,
    include_associated_revenue: bool,
) -> None:
    worksheet = workbook.create_sheet("Справка")
    worksheet.sheet_view.showGridLines = False
    worksheet.sheet_properties.tabColor = "8497B0"
    worksheet.freeze_panes = "A2"
    worksheet.column_dimensions["A"].width = 22
    worksheet.column_dimensions["B"].width = 42
    worksheet.column_dimensions["C"].width = 72
    worksheet.column_dimensions["D"].width = 72
    worksheet.row_dimensions[1].height = 32

    headers = ["Раздел", "Столбец", "Что означает", "Формула / важная особенность"]
    worksheet.append([_header_cell(worksheet, value) for value in headers])
    rows = _column_help_rows(cost_includes_vat, include_associated_revenue)
    for row_number, values in enumerate(rows, start=2):
        is_channel_role = values[0] == "Роль каналов"
        worksheet.row_dimensions[row_number].height = 60 if is_channel_role else 42
        cells = []
        for index, value in enumerate(values):
            cell = _body_cell(worksheet, value, "help")
            cell.alignment = Alignment(
                horizontal="left", vertical="top",
                wrap_text=index >= 2 or (is_channel_role and index == 1)
            )
            cells.append(cell)
        worksheet.append(cells)
    worksheet.auto_filter.ref = f"A1:D{len(rows) + 1}"


def _add_qa_sheet(workbook: Workbook, qa: dict[str, object]) -> None:
    worksheet = workbook.create_sheet("Сводка")
    worksheet.sheet_view.showGridLines = False
    worksheet.sheet_properties.tabColor = "1F4E78"
    worksheet.column_dimensions["A"].width = 42
    worksheet.column_dimensions["B"].width = 32
    worksheet.column_dimensions["C"].width = 22
    worksheet.column_dimensions["D"].width = 20
    worksheet.column_dimensions["E"].width = 14

    title = WriteOnlyCell(
        worksheet, value="Ассоциированные конверсии: Директ + Метрика"
    )
    title.font = TITLE_FONT
    worksheet.append([title])
    worksheet.append([])

    cohort = qa.get("conversion_cohort", {})
    parameters = [
        ["Параметр", "Значение"],
        ["Период", f"{cohort.get('date_from', '')} — {cohort.get('date_to', '')}"],
        ["Счётчики Метрики", ", ".join(map(str, qa.get("metrika_counter_ids", [])))],
        ["Цели", ", ".join(map(str, qa.get("goal_ids", [])))],
        ["Окно ассиста, дней", cohort.get("lookback_days")],
        [
            "Максимальный период истории",
            f"{cohort.get('association_date_from', '')} — "
            f"{cohort.get('association_date_to', '')}",
        ],
        [
            "Условие ассоциированности",
            "Целевой визит не из Директа; до него был Директ в индивидуальном окне",
        ],
        ["База лидов Метрики", "Уникальные визиты с целью"],
        [
            "Модель атрибуции Директа",
            _safe_value(
                cohort.get("direct_attribution_model"),
                "direct_attribution_model",
            ),
        ],
        [
            "Модель атрибуции Метрики",
            _safe_value(
                cohort.get("metrika_attribution_model"),
                "metrika_attribution_model",
            ),
        ],
        [
            "НДС в расходах Директа",
            "Включён" if cohort.get("direct_cost_includes_vat") else "Не включён",
        ],
        [
            "Ассоциированный доход",
            "Включён"
            if cohort.get("include_associated_revenue")
            else "Не включён",
        ],
        ["Источник расходов", "Yandex Direct Reports API"],
        ["Источник визитов", "Yandex Metrica Logs API"],
    ]
    for row_index, values in enumerate(parameters):
        if row_index == 0:
            worksheet.append([_header_cell(worksheet, str(value)) for value in values])
        else:
            worksheet.append(
                [_body_cell(worksheet, value, "parameter") for value in values]
            )

    worksheet.append([])
    section = WriteOnlyCell(worksheet, value="Загрузка данных")
    section.font = SECTION_FONT
    worksheet.append([section])
    loaded_labels = {
        "metrika_visits": "Визиты Метрики",
        "selected_goal_visit_rows": "Строки визитов с выбранными целями",
        "direct_rows": "Строки Директа",
        "conversion_touches": "Предыдущие визиты Директа в путях",
        "conversion_paths": "Уникальные агрегированные пути",
    }
    worksheet.append(
        [_header_cell(worksheet, "Показатель"), _header_cell(worksheet, "Строки")]
    )
    loaded = qa.get("loaded", {})
    for key, label in loaded_labels.items():
        worksheet.append(
            [
                _body_cell(worksheet, label, "metric"),
                _body_cell(worksheet, loaded.get(key), "visits"),
            ]
        )

    worksheet.append([])
    section = WriteOnlyCell(worksheet, value="Выбранные цели")
    section.font = SECTION_FONT
    worksheet.append([section])
    goal_headers = ["ID цели", "Визиты с целью", "Достижения цели"]
    if cohort.get("include_associated_revenue"):
        goal_headers.extend(["Доход цели", "Валюта"])
    worksheet.append(
        [_header_cell(worksheet, value) for value in goal_headers]
    )
    for goal_id, detail in qa.get("selected_goals_detail", {}).items():
        goal_row = [
            _body_cell(worksheet, goal_id, "goal_id"),
            _body_cell(worksheet, detail.get("goal_visits"), "visits"),
            _body_cell(worksheet, detail.get("goal_reaches"), "reaches"),
        ]
        if cohort.get("include_associated_revenue"):
            goal_row.extend(
                [
                    _body_cell(
                        worksheet, detail.get("goal_revenue"), "goal_revenue"
                    ),
                    _body_cell(
                        worksheet, detail.get("goal_currency"), "goal_currency"
                    ),
                ]
            )
        worksheet.append(goal_row)

    worksheet.append([])
    section = WriteOnlyCell(worksheet, value="Контрольные итоги")
    section.font = SECTION_FONT
    worksheet.append([section])
    totals = qa.get("combined_totals", {})
    total_rows = [
        ["Метрика", "Значение"],
        ["Показы Директа", totals.get("direct_impressions")],
        ["Клики Директа", totals.get("direct_clicks")],
        [
            "Расход Директа с НДС"
            if cohort.get("direct_cost_includes_vat")
            else "Расход Директа без НДС",
            totals.get("direct_cost"),
        ],
        ["Доход Директа", totals.get("direct_revenue")],
        ["ДРР Директа, %", totals.get("direct_drr_pct")],
        ["ROMI Директа, %", totals.get("direct_romi_pct")],
        ["Лиды Директа", totals.get("direct_api_converted_visits_sum")],
        ["Лиды Метрики", totals.get("unique_goal_visits")],
        ["Ассоциированные конверсии Метрики", totals.get("assisted_metrika_conversions_unique")],
        ["Всего конверсий Метрики", totals.get("metrika_conversions_with_assists")],
        ["Достижения целей", totals.get("goal_reaches")],
        ["Строки предыдущих визитов Директа", totals.get("conversion_touch_rows_non_additive")],
        ["Линейный кредит ассистов по источникам", totals.get("linear_conversion_credit_total")],
    ]
    if cohort.get("include_associated_revenue"):
        total_rows.append(
            [
                "Ассоциированный доход Метрики",
                totals.get("assisted_metrika_revenue_unique"),
            ]
        )
    for row_index, values in enumerate(total_rows):
        if row_index == 0:
            worksheet.append([_header_cell(worksheet, str(value)) for value in values])
        else:
            value_column = {
                "ДРР Директа, %": "drr_direct_pct",
                "ROMI Директа, %": "romi_direct_pct",
            }.get(values[0], "value")
            worksheet.append(
                [
                    _body_cell(worksheet, values[0], "metric"),
                    _body_cell(worksheet, values[1], value_column),
                ]
            )

    worksheet.append([])
    section = WriteOnlyCell(worksheet, value="Примечания")
    section.font = SECTION_FONT
    worksheet.append([section])
    for note in qa.get("notes", []):
        worksheet.append([_body_cell(worksheet, str(note), "note")])


def export_excel_workbook(
    connection: sqlite3.Connection,
    output_path: Path,
    qa: dict[str, object],
    channel_assists_sql: str,
    conversion_touches_sql: str,
) -> dict[str, object]:
    """Stream every report into one formatted XLSX workbook."""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook(write_only=True)
    _add_qa_sheet(workbook, qa)
    cohort = qa.get("conversion_cohort", {})
    cost_includes_vat = bool(cohort.get("direct_cost_includes_vat"))
    include_associated_revenue = bool(
        cohort.get("include_associated_revenue")
    )
    _add_help_sheet(
        workbook,
        cost_includes_vat,
        include_associated_revenue,
    )

    sheet_counts: dict[str, int] = {}
    for sheet_name, sql in SHEET_SPECS:
        sheet_counts[sheet_name] = _add_query_sheet(
            workbook, connection, sheet_name, sql, cost_includes_vat
        )
    sheet_counts["Ассисты каналов"] = _add_query_sheet(
        workbook,
        connection,
        "Ассисты каналов",
        channel_assists_sql,
        cost_includes_vat,
    )
    sheet_counts["Цепочки конверсий"] = _add_query_sheet(
        workbook,
        connection,
        "Цепочки конверсий",
        conversion_touches_sql,
        cost_includes_vat,
    )

    temporary_path = output_path.with_name(f".{output_path.name}.tmp")
    workbook.save(temporary_path)
    temporary_path.replace(output_path)
    return {
        "file": output_path.name,
        "sheets": ["Сводка", "Справка", *sheet_counts.keys()],
        "rows": sheet_counts,
    }
